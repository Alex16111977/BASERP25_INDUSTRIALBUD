# -*- coding: utf-8 -*-
"""Звірка реєстрів АВР книги «Звітність ВИРОБНИЦТВО» з актами фінзвіту виробництва (з розшифровкою).

Акти фінзвіту: рядки розшифровки під договором зі статусом «оригинал получен» (підписані, кол. E) /
«не подписан» (подані, проведені без оригіналу, кол. H) / «не проведён» (кол. H); примітка на B — дані реєстру АВР
з 1С (№ акта замовника, подання на ВТВ, статус). Реєстри книги — листи «РеєстрАВР_<об'єкт>» виробництва (формати
34 і 15 колонок, `knowledge_РеєстрАВР/02`); листи з префіксом «1» — вивантаження з 1С, у звірку не беруться.
Зіставлення: хвіст номера акта + сума ±0,01 («Роботи з ПДВ», інакше «Всього»), далі — № акта замовника з примітки.

Запуск: python -B zvirka.py <фінзвіт.xlsx> <книга.xlsx> <результат.xlsx>  (лише читання вхідних файлів)."""
import sys, os, re, datetime as dt, tempfile, zipfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# лист книги → об'єкт фінзвіту (назва до « (»)
LISTY = {
    "РеєстрАВР_Чорноморськ": "МК Чорноморськ",
    "РеєстрАВР_Кривий Ріг": "МК Кривий ріг",
    "РеєстрАВР_Кам'янець": "МК Каменец Цемент",
    "РеєстрАВР_Глобино)": "МК Глобино",
    "РеєстрАВР_Оріль": "МД МХП ОРІЛЬ",
    "РеєстрАВР_IRC": "МД IRC 2026",
}
# рамкові договори: у 1С зведені акти з іншою нумерацією й розбивкою — звіряються підсумками об'єкта
LISTY_PIDSUMKY = {"РеєстрАВР_ООН": "МД ООН 2025"}
STATUSY_1C = ("оригинал получен", "не подписан", "не проведён")


def wb_load(path, data_only=True):
    try:
        return openpyxl.load_workbook(path, data_only=data_only)
    except KeyError:  # 1С кладе xl/SharedStrings.xml з великої S
        out = os.path.join(tempfile.gettempdir(), "_fix_" + os.path.basename(path))
        with zipfile.ZipFile(path) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
            for it in zi.infolist():
                data, name = zi.read(it.filename), it.filename
                if name == "xl/SharedStrings.xml":
                    name = "xl/sharedStrings.xml"
                if name == "[Content_Types].xml" or name.endswith(".rels"):
                    data = data.replace(b"SharedStrings.xml", b"sharedStrings.xml")
                zo.writestr(name, data)
        return openpyxl.load_workbook(out, data_only=data_only)


def num(x):
    return float(x) if isinstance(x, (int, float)) else 0.0


def fmt(x):
    return f"{x:,.2f}".replace(",", " ").replace(".", ",") if x is not None else ""


def d2s(d):
    return d.strftime("%d.%m.%Y") if isinstance(d, (dt.date, dt.datetime)) else (str(d) if d else "")


def to_date(x):
    if isinstance(x, dt.datetime):
        return x.date()
    if isinstance(x, dt.date):
        return x
    m = re.search(r"(\d\d)\.(\d\d)\.(\d{4})(?!\d)", str(x or ""))
    if m and 2023 <= int(m.group(3)) <= 2027:
        try:
            return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def tail(n):
    m = re.search(r"(\d+)\s*$", str(n or "").strip())
    return int(m.group(1)) if m else None


def cat_book(s):
    s = (s or "").lower().replace("прдані", "подані")
    if "підписано" in s:
        return "підписано"
    if "перевірці" in s or "подані" in s or "подано" in s:
        return "подано"
    return "не подано" if s.strip() else "без статусу"


def cat_1c(s):
    return {"оригинал получен": "підписано", "не подписан": "подано", "не проведён": "не проведено"}.get(s, s)


# ---------------- фінзвіт ----------------
def read_finzvit(path):
    wv, wf = wb_load(path, True), wb_load(path, False)
    ws, wsf = wv.worksheets[0], wf.worksheets[0]
    akty, obj, dog = [], None, None
    for r in range(6, ws.max_row + 1):
        a = ws["A%d" % r].value
        c = wsf["C%d" % r].value
        if a and isinstance(c, str) and re.fullmatch(r"=C%d" % (r + 1), c):
            obj = str(a).split(" (")[0].strip()
            dog = None
            continue
        if obj is None or not a:
            continue
        if str(a).strip() == "Итого":
            obj = dog = None
            continue
        k, e_f = wsf["K%d" % r].value, wsf["E%d" % r].value
        if (isinstance(k, str) and "MAX(0" in k.upper()) or (isinstance(e_f, str) and e_f.upper().startswith("=SUBTOTAL(9,E")):
            dog = str(a).strip()  # рядок договору (у рамкових договорів K без формули — за SUBTOTAL колонки E)
            continue
        b = str(ws["B%d" % r].value or "")
        if dog is None or str(a).startswith("Поступление") or not b.startswith(STATUSY_1C):
            continue
        m = re.search(r"(\S+-\d+) от (\d\d\.\d\d\.\d{4})", str(a))
        komm = wsf["B%d" % r].comment.text if wsf["B%d" % r].comment else ""
        mz = re.search(r"АВР (\S+)", komm)
        mv = re.search(r"подано на ВТВ (\d\d\.\d\d\.\d{4})", komm)
        ms = re.search(r"! статус АВР «([^»]+)»", komm)
        e, h = num(ws["E%d" % r].value), num(ws["H%d" % r].value)
        akty.append({"obj": obj, "dog": dog, "row": r, "nomer": m.group(1) if m else str(a), "tail": tail(m.group(1)) if m else None,
                     "data": to_date(m.group(2)) if m else None, "status": b.split(" ")[0] + (" " + b.split(" ")[1] if " " in b else ""),
                     "status_full": b, "kat": cat_1c(b if b in STATUSY_1C else ("оригинал получен" if b.startswith("оригинал") else b)),
                     "suma": e or h, "E": e, "H": h, "zam": mz.group(1).rstrip("·,;") if mz else "",
                     "vtv": to_date(mv.group(1)) if mv else None, "st_reestr_1c": ms.group(1) if ms else "",
                     "oplata": str(ws["N%d" % r].value or "").split("\n")[0], "komm": komm})
    return akty


# ---------------- книга ----------------
def read_book(path):
    wb = wb_load(path, True)
    rows = []
    for name in wb.sheetnames:
        if not name.startswith("РеєстрАВР_"):
            continue
        ws = wb[name]
        hdr_r, cols = None, {}
        for r in range(1, 8):
            vals = {c: str(ws.cell(r, c).value or "").strip() for c in range(1, ws.max_column + 1)}
            if "№ акту" in vals.values():
                hdr_r = r
                for c, v in vals.items():
                    if v and v not in cols:
                        cols[v] = c
                break
        if not hdr_r:
            continue
        c_n, c_st, c_d = cols["№ акту"], cols["Статус"], cols.get("Дата") or cols.get("Дата ")
        c_d = c_d or [c for v, c in cols.items() if v.startswith("Дата")][0]
        c_vs = [c for v, c in cols.items() if v.startswith("Всього, грн")][0]
        c_rob = [c for v, c in cols.items() if v.startswith("Роботи")][0]
        dc = ""
        for r in range(hdr_r + 1, ws.max_row + 1):
            a = str(ws.cell(r, 1).value or "").strip()
            if a.lower().startswith(("всього", "різниця")) or str(ws.cell(r, 2).value or "").lower().startswith(("всього", "різниця")):
                break
            n = ws.cell(r, c_n).value
            if a and a not in ("0",):
                dc = a
            t = tail(n) if n is not None and str(n).strip() not in ("", "None") else None
            if t is None:
                continue
            vtv = to_date(ws.cell(r, c_d + 3).value) or to_date(ws.cell(r, c_d + 2).value)
            vtv_raw = str(ws.cell(r, c_d + 2).value or "").strip()
            rows.append({"list": name, "row": r, "dc": dc, "nomer": str(n).strip(), "tail": t, "vsogo": num(ws.cell(r, c_vs).value),
                         "roboty": num(ws.cell(r, c_rob).value), "status": str(ws.cell(r, c_st).value or "").strip(),
                         "kat": cat_book(str(ws.cell(r, c_st).value or "")), "data": to_date(ws.cell(r, c_d).value),
                         "zam": str(ws.cell(r, c_d + 1).value or "").strip(), "vtv": vtv, "vtv_raw": vtv_raw,
                         "oplata": ws.cell(r, c_d + 4).value})
    return rows


# ---------------- зіставлення ----------------
def zistavyty(akty, book):
    pary, vilni_b = [], {id(b): b for b in book}
    for a in akty:
        list_ = [k for k, v in LISTY.items() if v == a["obj"]]
        kand = [b for b in vilni_b.values() if b["list"] in list_ and b["tail"] == a["tail"]]
        best = None
        for b in kand:
            if abs(b["roboty"] - a["suma"]) < 0.01 or abs(b["vsogo"] - a["suma"]) < 0.01:
                best = b
                break
        if best is None and a["zam"]:
            for b in vilni_b.values():
                if b["list"] in list_ and b["zam"] == a["zam"] and (abs(b["roboty"] - a["suma"]) < 0.01 or abs(b["vsogo"] - a["suma"]) < 0.01):
                    best = b
                    break
        if best is None and len(kand) == 1:
            best = kand[0]  # той самий номер, сума інша
        if best is None:
            same = [b for b in vilni_b.values() if b["list"] in list_ and (abs(b["roboty"] - a["suma"]) < 0.01 or abs(b["vsogo"] - a["suma"]) < 0.01)]
            if len(same) == 1:
                best = same[0]  # номер інший, сума єдина
        if best is not None:
            del vilni_b[id(best)]
        pary.append((a, best))
    return pary, list(vilni_b.values())


def rozbizhnosti(a, b):
    out = []
    if b is None:
        out.append(("Немає в реєстрі книги", "акт є у фінзвіті (%s, %s), у реєстрі виробництва рядка немає" % (cat_1c_text(a), fmt(a["suma"]))))
        return out
    sb = b["roboty"] or b["vsogo"]
    if b["tail"] != a["tail"]:
        out.append(("№ акта", "у реєстрі «%s», у 1С %s (зіставлено за сумою)" % (b["nomer"], a["nomer"])))
    if abs(b["roboty"] - a["suma"]) >= 0.01 and abs(b["vsogo"] - a["suma"]) >= 0.01:
        out.append(("Сума", "реєстр %s ≠ акт 1С %s (різниця %s)" % (fmt(sb), fmt(a["suma"]), fmt(sb - a["suma"]))))
    if b["kat"] == "підписано" and a["kat"] != "підписано":
        out.append(("Статус", "у реєстрі «%s», у 1С — %s (немає відмітки «Оригінал»)" % (b["status"], cat_1c_text(a))))
    elif a["kat"] == "підписано" and b["kat"] != "підписано":
        out.append(("Статус", "у 1С оригінал отримано, у реєстрі — «%s»" % b["status"]))
    elif a["kat"] == "не проведено" and b["kat"] == "подано":
        out.append(("Не проведено", "у реєстрі «%s», у 1С акт не проведено" % b["status"]))
    elif b["kat"] in ("не подано", "без статусу") and a["kat"] != "підписано":
        out.append(("Статус", "у реєстрі «%s», у 1С — %s" % (b["status"] or "статус порожній", cat_1c_text(a))))
    if a["zam"] and b["zam"] and a["zam"] != b["zam"]:
        out.append(("№ акта замовника", "реєстр «%s», 1С «%s»" % (b["zam"], a["zam"])))
    if a["vtv"] and not b["vtv"] and b.get("vtv_raw"):
        out.append(("Дата подання на ВТВ", "у реєстрі помилкова дата «%s», у 1С %s" % (b["vtv_raw"], d2s(a["vtv"]))))
    if a["vtv"] and b["vtv"] and a["vtv"] != b["vtv"]:
        out.append(("Дата подання на ВТВ", "реєстр %s, 1С %s" % (d2s(b["vtv"]), d2s(a["vtv"]))))
    return out


def cat_1c_text(a):
    return {"підписано": "оригінал отримано", "подано": "проведено, не підписано", "не проведено": "не проведено"}.get(a["kat"], a["status_full"])


def main(fz, kn, out):
    akty, book = read_finzvit(fz), read_book(kn)
    akty_obj = [a for a in akty if a["obj"] not in LISTY_PIDSUMKY.values()]
    pary, zalyshok = zistavyty(akty_obj, book)
    vyb = openpyxl.Workbook()
    ws = vyb.active
    ws.title = "Розбіжності"
    zag = ["Об'єкт", "Акт 1С", "Дата акта", "Статус у 1С", "Сума акта 1С", "Лист книги / рядок", "№ акта у реєстрі",
           "Статус у реєстрі", "Сума в реєстрі", "Що не так", "Пояснення"]
    ws.append(zag)
    n_r = 0
    for a, b in pary:
        for vyd, txt in rozbizhnosti(a, b):
            n_r += 1
            ws.append([a["obj"], a["nomer"], d2s(a["data"]), cat_1c_text(a), a["suma"], ("%s / %d" % (b["list"], b["row"])) if b else "",
                       b["nomer"] if b else "", b["status"] if b else "", (b["roboty"] or b["vsogo"]) if b else None, vyd, txt])
    for b in zalyshok:
        if b["list"] in LISTY and b["kat"] != "підписано":
            n_r += 1
            ws.append([LISTY[b["list"]], "", "", "", None, "%s / %d" % (b["list"], b["row"]), b["nomer"], b["status"],
                       b["roboty"] or b["vsogo"], "Немає у фінзвіті", "рядок реєстру без акта у фінзвіті (номер/сума не знайдені)"])
    ws2 = vyb.create_sheet("Акти фінзвіту")
    ws2.append(["Об'єкт", "Договір", "Акт 1С", "Дата акта", "Статус у 1С", "Сума E (підписані)", "Сума H (подані/не проведені)",
                "№ акта зам. (1С)", "Подання на ВТВ (1С)", "Лист книги / рядок", "Статус у реєстрі", "Сума в реєстрі",
                "№ акта зам. (реєстр)", "Подання на ВТВ (реєстр)", "Оплата (реєстр)", "Оплата (фінзвіт)", "Результат"])
    for a, b in pary:
        rz = rozbizhnosti(a, b)
        ws2.append([a["obj"], a["dog"], a["nomer"], d2s(a["data"]), cat_1c_text(a), a["E"] or None, a["H"] or None, a["zam"], d2s(a["vtv"]),
                    ("%s / %d" % (b["list"], b["row"])) if b else "", b["status"] if b else "", (b["roboty"] or b["vsogo"]) if b else None,
                    b["zam"] if b else "", d2s(b["vtv"]) if b else "", (str(b["oplata"]) if b and b["oplata"] is not None else ""),
                    a["oplata"], "; ".join(v for v, _ in rz) or "збігається"])
    ws3 = vyb.create_sheet("Інші листи книги")
    ws3.append(["Лист", "Рядків актів", "Пояснення"])
    from collections import Counter
    cnt = Counter(b["list"] for b in book)
    for l, n in sorted(cnt.items()):
        if l in LISTY_PIDSUMKY:
            ob = LISTY_PIDSUMKY[l]
            s1c = sum(a["suma"] for a in akty if a["obj"] == ob)
            sbk = sum((b["roboty"] or b["vsogo"]) for b in book if b["list"] == l)
            ws3.append([l, n, "рамковий договір: у 1С %d зведених актів на %s (усі «оригінал отримано»), у реєстрі %d рядків на %s; "
                        "нумерація й розбивка різні — поактна звірка неможлива" % (sum(1 for a in akty if a["obj"] == ob), fmt(s1c), n, fmt(sbk))])
        elif l not in LISTY:
            ws3.append([l, n, "об'єкта немає у фінзвіті — не звірявся"])
        else:
            st = [b for b in zalyshok if b["list"] == l and b["kat"] == "підписано"]
            if st:
                ws3.append([l, len(st), "підписані акти реєстру, яких немає в розшифровці фінзвіту (закриті договори поза фінзвітом): "
                            + ", ".join("%s (%s)" % (b["nomer"], fmt(b["roboty"] or b["vsogo"])) for b in st)])
    for w in (ws, ws2, ws3):
        for c in w[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D7E4BD")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        w.freeze_panes = "A2"
        for i, col in enumerate(w.columns, 1):
            w.column_dimensions[get_column_letter(i)].width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
        for row in w.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00"
                c.alignment = Alignment(wrap_text=True, vertical="top")
    vyb.save(out)
    return akty, book, pary, zalyshok, n_r


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    akty, book, pary, zal, n = main(*sys.argv[1:4])
    print("актів у фінзвіті:", len(akty), "| рядків актів у реєстрах книги:", len(book), "| розбіжностей:", n)
    for a, b in pary:
        rz = rozbizhnosti(a, b)
        if rz or a["kat"] != "підписано":
            print("%-18s %-14s %-11s %-24s %14s | %-22s %-34s | %s" % (
                a["obj"][:18], a["nomer"], d2s(a["data"]), cat_1c_text(a), fmt(a["suma"]), (b["list"][10:] + "/" + str(b["row"])) if b else "—",
                (b["status"] if b else "")[:34], "; ".join("%s: %s" % x for x in rz) or "OK"))
    for b in zal:
        if b["list"] in LISTY and b["kat"] != "підписано":
            print("ЛИШЕ В КНИЗІ:", b["list"], b["row"], b["nomer"], b["status"], fmt(b["roboty"] or b["vsogo"]))
