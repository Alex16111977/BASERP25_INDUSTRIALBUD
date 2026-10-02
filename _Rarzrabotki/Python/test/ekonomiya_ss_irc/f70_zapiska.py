# -*- coding: utf-8 -*-
"""Аналітична записка директору (HTML, укр.) + додаток xlsx (укр.) за результатами дослідження
«економія матеріалів проти СС, МД IRC 2026». Усі числа — з out/*.json (f40/f41/f60) і файлів-звітів; суми таблиць
сходяться до копійки за побудовою (залишок округлення — методом найбільшого залишку). Бази не чіпає.
Версія 2 (02.10.2026): правки за перевіркою workflow verify-director-note-irc (74 підтверджені зауваження)."""
import os, sys, json, re, html
sys.path.insert(0, os.path.dirname(__file__))
from common import *
from collections import defaultdict, OrderedDict

L = lambda n: json.load(open(os.path.join(OUT, n), encoding="utf-8"))
SE_KOM = json.load(open(os.path.join(OUT, "ss_erp.json"), encoding="utf-8"))["kom"]
A = L("analysis.json")["rows"]; C = L("core.json"); F = L("final.json"); RK = L("rk_irc.json"); SB = L("ss_buh.json")
ON = {k: defaultdict(Decimal, {m: D(v) for m, v in d.items()}) for k, d in C["on"].items()}
Z = Decimal(0)
q2 = lambda x: D(x).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
dv = lambda a, b: a / b if b else Z
H15 = [f"15-{i}" for i in range(1, 7)]; H30 = ["30-1", "30-2"]
LOG = []

def isnum(s):
    try: Decimal(s); return True
    except Exception: return False
R = {k: {f: (D(v) if isinstance(v, str) and isnum(v) and f not in ("k", "name") else v) for f, v in r.items()} for k, r in A.items()}
classify = classify_row

def alloc(parts, target, tag):
    """Округлити складові так, щоб сума = target (вже округлений): залишок — на найбільший залишок округлення."""
    rd = OrderedDict((k, q2(v)) for k, v in parts.items())
    diff = target - sum(rd.values())
    if diff:
        big = max(parts, key=lambda k: (parts[k] - rd[k]) * (1 if diff > 0 else -1))
        rd[big] += diff
        LOG.append(f"округлення {tag}: {diff} → «{big}»")
    return rd

ABBR = {"мшп": "МШП", "пвл": "ПВЛ", "осб": "ОСБ", "цсп": "ЦСП", "пвх": "ПВХ", "гк": "ГК", "led": "LED"}
def nm(s):
    s = " ".join(str(s).split())
    s = " ".join(ABBR.get(wd.casefold(), wd) for wd in s.split(" "))
    return s[:1].upper() + s[1:]

# ---------- складові економії за витратою (15 м, 30 м)
CATN = defaultdict(list)            # (група, категорія) -> [(|Δ|, назва)]
def parts(g, houses):
    p = defaultdict(Decimal)
    for r in R.values():
        o = ON[r["k"]]; c = classify(r, g); dl = r[f"{g}_delta"]
        if c == "сопоставимо":
            ns = o[f"nr|{houses[0]}|ПланГрн"]
            z = sum(1 for h in houses if o[f"c|{h}|q"] == 0)
            pz = (Decimal(z) * ns) if 0 < z < len(houses) else Z     # норма будинків без списання (грн)
            p["price"] += r[f"{g}_pe"]; p["qty"] += r[f"{g}_qe"] + pz; p["partial"] -= pz
            if pz: CATN[(g, "partial")].append((pz, r["name"]))
        elif c == "поза кошторисом":
            p["out"] += dl; CATN[(g, "out")].append((abs(dl), r["name"]))
        elif c == "не использовано":
            p["notused"] += dl; CATN[(g, "notused")].append((abs(dl), r["name"]))
        elif c in ("единица/состав", "несопоставимо"):
            z = sum(1 for h in houses if o[f"c|{h}|q"] == 0)
            pz = (Decimal(z) * o[f"nr|{houses[0]}|ПланГрн"]) if 0 < z < len(houses) else Z
            if dl + pz > 0:            # перерасход на домах со списанием — подтверждён; норма домов без списания — нет
                p["comp"] += dl + pz; CATN[(g, "comp")].append((dl + pz, r["name"]))
                if pz:
                    p["partial"] -= pz; p["partial_nd"] -= pz; CATN[(g, "partial")].append((pz, r["name"]))
            else:
                p["incomp"] += dl; CATN[(g, "incomp")].append((abs(dl), r["name"]))
        else:
            p["other"] += dl
    return p

def names(key, n=4, groups=("d15", "d30")):
    seen = []; lst = sorted([x for g in groups for x in CATN[(g, key)]], reverse=True)
    for _, nme in lst:
        nme = nm(nme).lower() if nme.lower() not in ("мшп",) else "МШП"
        nme = nm(nme) if nme[:1].isupper() else nme
        nme = {"грунт": "ґрунт"}.get(nme, nme)
        if nme not in seen: seen.append(nme)
    return seen[:n]

P15, P30 = parts("d15", H15), parts("d30", H30)
assert P15["other"] == 0 and P30["other"] == 0
PND = {"15": q2(P15["partial_nd"]), "30": q2(P30["partial_nd"])}
for P_ in (P15, P30): P_.pop("partial_nd", None)
left = {h: sum(o[f"stkh|{h}|gt"] for o in ON.values()) for h in H15 + H30 + ["30-3"]}
cons = {h: sum(o[f"c|{h}|gt"] for o in ON.values()) for h in H15 + H30 + ["30-3"]}
col = {}
for g, P, hs in (("15", P15, H15), ("30", P30, H30)):
    delta = sum(P.values()); lf = sum(left[h] for h in hs)
    D_r = q2(delta); LF_r = q2(lf)
    LB_r = q2(P["price"] + P["qty"] + P["out"] + P["comp"] + lf)          # мінус = економія
    S1_r = LB_r - LF_r; S2_r = D_r - S1_r
    g1 = alloc(OrderedDict(price=P["price"], qty=P["qty"], out=P["out"], comp=P["comp"]), S1_r, f"{g} підтверджене")
    g2 = alloc(OrderedDict(notused=P["notused"], partial=P["partial"], incomp=P["incomp"]), S2_r, f"{g} непідтверджене")
    col[g] = dict(**g1, **g2, S1=S1_r, S2=S2_r, delta=D_r, left=LF_r, LB=LB_r, PROB=D_r + LF_r,
                  norm=q2(sum(o[f"nr|{hs[0]}|ПланГрн"] for o in ON.values()) * len(hs)), cons=q2(sum(cons[h] for h in hs)))
tot = {k: col["15"][k] + col["30"][k] for k in col["15"]}
assert tot["delta"] == q2(-D(F["d15"]) - D(F["d30"])), (tot["delta"], F["d15"], F["d30"])
assert -tot["LB"] == q2(F["LB"]), (tot["LB"], F["LB"])
assert tot["norm"] == D("6247401.54") and tot["cons"] - tot["norm"] == tot["delta"]
N8 = tot["norm"]
assert -tot["PROB"] == q2(F["probable"])

# ---------- щотижневий звіт (як у xlsx 30.09) і детальний варіант (irc_3009.json)
B = os.path.join(ROOT, r"_Rarzrabotki\Письмо Директору по отчетам на производство")
wsx = load_xlsx(os.path.join(B, r"30092026\IRC 2026\План-факт виконання СС на 30-09-2026.xlsx")).worksheets[0]
rep = None
for row in wsx.iter_rows(values_only=True):
    if row and isinstance(row[0], str) and row[0].strip() in ("МД IRC 2026", "МД IRС 2026"):
        rep = [c for c in row if c is not None]; break
PLAN13, PNF, DEV = q2(rep[5]), q2(rep[6]), q2(rep[7])
assert (PLAN13, PNF, DEV) == (D("11296725.24"), D("10902355.49"), D("-394369.75"))
det = json.load(open(os.path.join(ROOT, r"_Rarzrabotki\Python\test\pismo_2026_09\irc_3009.json"), encoding="utf-8"))
num = lambda s: D(str(s).replace(" ", "").replace(" ", "").replace(",", "."))
dr = next(x for x in det if x["cells"][0].strip() in ("МД IRC 2026", "МД IRС 2026"))
PUR, ETC = num(dr["cells"][6]), num(dr["cells"][8])
assert PUR + ETC == PNF and num(dr["cells"][9]) == PNF
NOTB = PUR - tot["cons"]                                  # закуплене, але не списане на 8 будинків
N5 = q2(D("1009864.74") * 5)
REST = NOTB + ETC - N5
assert col["15"]["delta"] + col["30"]["delta"] + REST == DEV, "розклад відхилення звіту не сходиться"
STOCK_P = q2(sum(o["stk|P|gt"] for o in ON.values()))
MSP_S = q2(sum(o["stk|MSP|gt"] for o in ON.values())); X_S = q2(sum(o["stk|X|gt"] for o in ON.values()))
H303 = q2(cons["30-3"]) + q2(left["30-3"])
OTHER = NOTB - tot["left"] - H303 - STOCK_P
RESID = OTHER - MSP_S - X_S
DEF3, DEF47 = q2(F["deficit303"]), q2(F["deficit47"]); DEF = DEF3 + DEF47; OVER = ETC - DEF

# ---------- звіт за нормами (xlsx як є)
wn = load_xlsx(os.path.join(B, r"Отчет\Списання IRC.xlsx")).worksheets[0]
NR_FACT = q2(wn["O13"].value) + q2(wn["O21"].value) + q2(wn["O22"].value)
NR_DEV = q2(wn["Q13"].value) + q2(wn["Q21"].value) + q2(wn["Q22"].value)
NR_EK = q2(wn["M13"].value) + q2(wn["M21"].value) + q2(wn["M22"].value)
NR_EK30 = q2(wn["M20"].value)
NR_EK3_7 = q2(wn["M23"].value) + sum(q2(wn[f"M{i}"].value) for i in range(24, 28))
assert NR_DEV == NR_FACT - N8
NR_O = {h: q2(wn[c].value) for h, c in zip(H15 + H30, ["O14", "O15", "O16", "O17", "O18", "O19", "O21", "O22"])}
assert sum(NR_O.values()) == NR_FACT
VAT_UP, VAT_DN = D("69856.18"), D("152034.53"); VAT_NET = VAT_DN - VAT_UP; SEMI = D("11879.84")   # f40 / details.txt

# ---------- бюджет (книга, аркуш «МД IRC 2026», рядок 106)
import openpyxl
BOOK = os.path.join(ROOT, r"_Rarzrabotki\Бюджет прогноз\Бюджет октябрь 2026\Принятый бюджет\Бюджет_Жовтень26-Грудень26.xlsx")
def row106(data_only):
    wb_ = openpyxl.load_workbook(BOOK, read_only=True, data_only=data_only)["МД IRC 2026"]
    return {c.coordinate: c.value for row in wb_.iter_rows(min_row=106, max_row=106) for c in row if c.value is not None}
cv, cf = row106(True), row106(False)
EC_EK, EC_PAID, EC_BUD, EC_TOT = q2(cv["U106"]), q2(cv["R106"]), q2(cv["S106"]), q2(cv["T106"])
m = re.match(r"=([\d.]+)\+([\d.]+)\*0\.8\+([\d.]+)\*0\.8", cf["B106"].replace(" ", ""))
EC_DEBT = D(m.group(1)); EC_NEW = D(m.group(2)) + D(m.group(3))
assert EC_DEBT + EC_NEW == EC_BUD and EC_PAID + EC_BUD == EC_TOT
LIM_D = q2(PLAN13 - D(str(cv["Q106"])))
UNPAID = PUR - EC_PAID; DEBT_X = EC_DEBT - UNPAID; NEW_D = ETC - EC_NEW
assert -DEV + NEW_D - DEBT_X - LIM_D == EC_EK, (-DEV + NEW_D - DEBT_X - LIM_D, EC_EK)

# ---------- форматування
NB = " "
def n(x, sign=False):
    x = q2(x); s = f"{abs(x):,.2f}".replace(",", NB).replace(".", ",")
    if sign: return ("−" if x < 0 else ("+" if x > 0 else "")) + s
    return ("−" if x < 0 else "") + s
def pct(x, d=1):
    q = Decimal("0.1") if d == 1 else Decimal("1")
    v = D(x).quantize(q, rounding=ROUND_HALF_UP)
    return (f"{v:.1f}".replace(".", ",") if d == 1 else f"{v:.0f}").replace("-", "−") + NB + "%"
def pc(a, b, d=1): return pct(dv(D(a), D(b)) * 100, d)
def cls(x): return "good" if x < 0 else ("bad" if x > 0 else "")
def td(x, sign=True, neutral=False): return f'<td class="num {"" if neutral else cls(q2(x))}">{n(x, sign)}</td>'
UN = {"м2": "м²", "куб": "м³", "м3": "м³", "пог.м": "пог. м", "Тонна": "т"}
def unit(r):
    u = [x.split(":", 1)[1].strip() for x in r["units"] if x.startswith("СС:")] or \
        [x.split(":", 1)[1].strip() for x in r["units"] if x.startswith("BUH:")]
    return UN.get(u[0], u[0]) if u and u[0] else "—"
def q2s(x): return n(x)
def fq(x):
    """Кількість: ціле — без копійок, інакше два знаки."""
    x = q2(x)
    return f"{int(x):,}".replace(",", NB) if x == x.to_integral_value() else n(x)
byname = lambda s: R[next(k for k in R if R[k]["name"] == s)]

# ---------- таблиця цін (15 м, зіставні, однакова одиниця)
okq = lambda r: D("0.5") <= dv(r["d15_cq"], r["d15_Nq"]) <= 2
pt = sorted([r for r in R.values() if classify(r, "d15") == "сопоставимо" and r["d15_pe"] < 0 and okq(r)], key=lambda r: r["d15_pe"])[:8]
price_rows = []
for r in pt:
    o = ON[r["k"]]; nq = o["nr|15-1|ПланКол"]; ns = o["nr|15-1|ПланГрн"]
    pss = ns / nq; pf = r["d15_cg"] / r["d15_cq"]
    price_rows.append((f"{nm(r['name'])}, {unit(r)}", pss, pf, pss - pf, dv(pss - pf, pss) * 100, r["d15_cq"], dv(r["d15_cq"], nq * 6) * 100, -q2(r["d15_pe"])))
PT_SUM = sum(x[7] for x in price_rows)
dear = sorted([r for r in R.values() if classify(r, "d15") == "сопоставимо" and r["d15_pe"] > 0 and okq(r)], key=lambda r: -r["d15_pe"])[:4]
gen = byname("генератор"); GEN_PE = -(q2(gen["d15_pe"]) + q2(gen["d30_pe"]))
GEN15_P = gen["d15_cg"] / gen["d15_cq"]; GEN30_P = gen["d30_cg"] / gen["d30_cq"]
GEN_PUR_P = gen["pu_g"] / gen["pu_q"]
pn, cs, gr = byname("Профнастил покрівельний"), byname("Плита ЦСП"), byname("Грунт")
plint = byname("Плінтус")
orig_pl = [x for x in SB["orig"] if " 15 м" in x["СС"] and norm_name(x["НомСС"]) == "Плінтус"][0]
kom_pl = [x for x in SB["kom"] if " 15 м" in x["СС"] and norm_name(x["НомСС"]) == "Плінтус"][0]
PL_TQ, PL_TS = D(str(kom_pl["Кол"])), D(str(kom_pl["Сумма"])); PL_OQ = D(str(orig_pl["КолОриг"] or orig_pl["Кол"])); PL_OS = D(str(orig_pl["Сумма"]))
assert PL_TS == PL_OS
ex = pn; exo = ON[ex["k"]]
l30_2 = {C["names"][k]: (D(d.get("stkh|30-2|gt", "0")), D(d.get("stkh|30-2|q", "0"))) for k, d in C["on"].items()}
CSP2, OSB2 = l30_2["Плита ЦСП"], l30_2["Плита ОСБ"]
GEN4 = D(C["on"][gen["k"]].get("stkh|15-4|gt", "0"))

# ---------- додаткові величини (друга перевірка)
SURPLUS = q2(sum(R[k]["P_left_after"] for k in R))                       # склад проєкту понад норму всіх незавершених
NONORM_P = q2(sum(ON[k]["stk|P|gt"] for k in R if R[k]["n15_s"] == 0 and R[k]["n30_s"] == 0))
COVER = STOCK_P - SURPLUS
OWN303 = q2(sum(R[k]["e303_from_own"] for k in R))
OWN_SURP = q2(left["30-3"]) - OWN303
FB37 = q2(cons["30-3"]) + OWN303 + COVER + DEF                             # норма 30 м №3–№7 за цінами запасу + нестача за ціною СС
ECON37 = N5 - FB37
REST2 = NOTB - tot["left"] + ETC - N5                                      # решта прогнозу звіту
MISC = OTHER + OWN_SURP
assert REST2 == OVER + SURPLUS - ECON37 + MISC, (REST2, OVER + SURPLUS - ECON37 + MISC)
assert tot["PROB"] + REST2 == DEV
P15_LIM = q2(D(str(cv["O106"])) - D("704612.01") * 6); P30_LIM = q2(D(str(cv["P106"])) - D("1009864.74") * 7)
assert P15_LIM + P30_LIM == -LIM_D, (P15_LIM, P30_LIM, LIM_D)
hk = byname("Гайка"); sh = byname("шайба"); ds = byname("Обрізні диски")
GK15 = -q2(hk["d15_delta"]); SH = -q2(sh["d15_delta"] + sh["d30_delta"]); GR30 = -q2(gr["d30_delta"]); DS15 = -q2(ds["d15_delta"])
INC_PARTS = GK15 + SH + GR30 + DS15
m48 = {("15" if " 15 м" in x["СС"] else "30"): D(str(x["Сумма"])) / D(str(x["Кол"])) for x in SE_KOM if "М48" in (x["НомСС"] or "") and x["ОН"] == "Гайка"}
sh12 = [D(str(x["Сумма"])) / D(str(x["Кол"])) for x in SE_KOM if (x["НомСС"] or "").lower().startswith("шайба м12")][0]
Q15_REST = -(col["15"]["notused"] + col["15"]["partial"])
INCOMP_TXT = -tot["incomp"]
ALIAS = [("d15", "Дозатор рідкого мила", "Диспенсер рідкого мила"), ("d15", "Дозатор рушників", "Диспенсер рушників"),
         ("d30", "кріплення ринви", "Держак ринви"), ("d30", "Кутики сталеві", "Кутики (металопрокат)")]
AL_NU = -sum(q2(byname(a)[f"{g}_delta"]) for g, a, b in ALIAS)
AL_OUT = sum(q2(byname(b)[f"{g}_delta"]) for g, a, b in ALIAS)
LB_AL = -tot["LB"] + AL_NU
GEN_LB = -tot["LB"] - GEN_PE

# ---------- HTML
H = []; w = H.append
TITLE = "МД IRC 2026 — економія матеріалів проти кошторису СС станом на 02.10.2026: аналітична записка"
w(f"""<!DOCTYPE html>
<html lang="uk">
<head>
<meta charset="utf-8">
<title>{TITLE}</title>
<style>
  :root{{--ink:#1f2933;--muted:#6b7683;--line:#dfe3e8;--good:#1b7f4b;--bad:#b3261e;--accent:#1f4e79;--accent-bg:#eef3f8}}
  *{{box-sizing:border-box}}
  body{{margin:0;padding:32px 28px 48px;background:#f4f6f8;color:var(--ink);font:15px/1.6 "Segoe UI",Tahoma,Arial,sans-serif}}
  .sheet{{max-width:960px;margin:0 auto;background:#fff;padding:40px 44px;border:1px solid var(--line);border-radius:6px;box-shadow:0 1px 4px rgba(0,0,0,.06)}}
  h1{{font-size:19px;margin:0 0 4px;color:var(--accent)}}
  .sub{{color:var(--muted);font-size:13px;margin:0 0 18px}}
  h2{{font-size:15px;margin:26px 0 8px;padding-bottom:5px;border-bottom:2px solid var(--accent-bg);color:var(--accent)}}
  h3{{font-size:14px;margin:16px 0 6px}}
  p{{margin:9px 0}}
  ul,ol{{margin:6px 0 6px 22px;padding:0}} li{{margin:5px 0}}
  table{{width:100%;border-collapse:collapse;margin:8px 0 4px;font-size:14px}}
  th,td{{padding:6px 10px;border-bottom:1px solid var(--line);text-align:right;vertical-align:top}}
  th:first-child,td:first-child{{text-align:left}}
  td.txt{{text-align:left}}
  thead th{{background:#f7f9fb;font-weight:600;color:var(--muted);font-size:13px;border-bottom:1px solid #c9d2da;line-height:1.3}}
  tr.total td{{font-weight:700;background:#f7f9fb;border-top:1px solid #c9d2da}}
  tr.grp td{{background:var(--accent-bg);color:var(--accent);font-weight:600;border-top:1px solid #c9d2da}}
  tr.sub td:first-child{{padding-left:26px;color:#3b4652}}
  tr.key td{{font-weight:700;border-top:2px solid var(--accent)}}
  .num{{font-variant-numeric:tabular-nums;white-space:nowrap}}
  .good{{color:var(--good);font-weight:600}}
  .bad{{color:var(--bad);font-weight:600}}
  .key{{border-left:4px solid var(--accent);background:var(--accent-bg);padding:10px 14px;margin:10px 0}}
  .guide{{background:#fafbfc;border:1px solid var(--line);border-radius:4px;padding:10px 16px;margin:8px 0}}
  .src{{color:var(--muted);font-size:12.5px;margin-top:22px;border-top:1px solid var(--line);padding-top:8px}}
  @page{{size:A4 portrait;margin:11mm}}
  @media print{{body{{background:#fff;padding:0;font-size:10.5pt}}.sheet{{border:0;box-shadow:none;max-width:none;padding:0}}
    h1{{font-size:14pt}}h2{{font-size:11.5pt;break-after:avoid}}table{{font-size:9.5pt;break-inside:avoid}}th,td{{padding:3px 6px}}.key,.guide{{break-inside:avoid}}}}
</style>
</head>
<body>
<div class="sheet">
""")
w("<p>Доброго дня,</p>")
w(f"<h1>{TITLE}</h1>")
w('<p class="sub">Закупівлі — по 30.09.2026; списання й залишки — бухоблік на 02.10.2026. Аналіз — 8 будинків, відвантажених замовнику: '
  '15 м №1–№6 і 30 м №1–№2. Будинки 30 м №3–№7 не завершені й до розрахунку економії не входять. Усі суми — з ПДВ.</p>')

# --- Коротко
w("<h2>Коротко</h2>")
w('<p><b>Економія матеріалів проти кошторису СС є, і вона більша за 3,5 %, які показує щотижневий план-факт.</b> '
  'Тут її пораховано не за закупівлями, а за тим, що фактично списано на кожен відвантажений будинок.</p>')
w(f'<div class="key">На 8 будинків списано матеріалів на <b class="num">{n(tot["cons"])}</b> грн проти норми кошторису СС '
  f'<b class="num">{n(N8)}</b> грн — на <b class="num good">{n(-tot["delta"])} грн ({pc(-tot["delta"], N8)})</b> менше.<br>'
  f'Імовірна економія — <b class="num good">{n(-tot["PROB"])} грн ({pc(-tot["PROB"], N8)})</b>: це різниця з нормою за мінусом '
  f'{n(tot["left"])} грн матеріалів, що ще лежать на складах цих будинків.<br>'
  f'Нижня межа з запасом, доведена документами, — <b class="num good">{n(-tot["LB"])} грн ({pc(-tot["LB"], N8)})</b>: 15 м №1–№6 — '
  f'{n(-col["15"]["LB"])} грн ({pc(-col["15"]["LB"], col["15"]["norm"])}), 30 м №1–№2 — {n(-col["30"]["LB"])} грн '
  f'({pc(-col["30"]["LB"], col["30"]["norm"])}); з неї {n(GEN_PE)} грн — генератор, потужність якого в документах закупівлі не вказана. '
  f'Решта імовірної економії, {n(-tot["S2"])} грн, потребує підтвердження.</div>')
w(f'<p>Економію дала ціна закупівлі: <b class="num">{n(-tot["price"])}</b> грн. За обсягом на 15 м витрата близька до норми: на будинках, '
  f'де кошторисні позиції списано, їх витрачено на {n(-col["15"]["qty"])} грн менше за норму, але додалися потрібні матеріали, яких у кошторисі '
  f'немає (тротуарна плитка, поручні, плінтус на стільницю, МШП), — на {n(col["15"]["out"])} грн. Ще {n(Q15_REST)} грн норми 15 м — '
  f'позиції, не списані на жоден будинок або не на всі; вони потребують підтвердження.</p>')
w(f'<p>Щотижневий план-факт показує економію {n(-DEV)} грн ({pc(-DEV, PLAN13)}): він враховує не списання на будинки, а закупівлі всіх '
  f'13 будинків — разом із {n(STOCK_P)} грн матеріалів на складі проєкту — і додає за ціною СС усе, що за кошторисом ще не закуплено '
  f'(«плановий розрахунок матеріалів», {n(ETC)} грн). За своєю формулою звіт порахований правильно (закупівлі збігаються до копійки, плановий розрахунок — з різницею 0,55 грн через округлення кількості у звіті), але як прогноз '
  f'закупівель він, за обліковими залишками, більший за потребу на {n(OVER)} грн, а економію за витратою на відвантажених будинках '
  f'показати не може.</p>')

# --- Таблиця 1
w("<h2>З чого складається економія на 8 будинках</h2>")
w('<p>Мінус — економія, плюс — перевитрата. Ця ж таблиця — у додатку, аркуш «З чого складається економія».</p>')
w('<table><thead><tr><th>Складова</th><th>15 м №1–№6</th><th>30 м №1–№2</th><th>Разом, грн</th></tr></thead><tbody>')
w(f'<tr><td>Норма кошторису СС</td><td class="num">{n(col["15"]["norm"])}</td><td class="num">{n(col["30"]["norm"])}</td><td class="num">{n(N8)}</td></tr>')
w(f'<tr><td>Списано фактично</td><td class="num">{n(col["15"]["cons"])}</td><td class="num">{n(col["30"]["cons"])}</td><td class="num">{n(tot["cons"])}</td></tr>')
LBL = OrderedDict(grp1="Складові, підтверджені документами закупівлі й списання",
                  price="Ціна закупівлі нижча за кошторис",
                  qty="Обсяг: списано менше чи більше за норму — на будинках, де позицію списано",
                  out="Списано матеріалів, яких немає в кошторисі",
                  comp="Інша одиниця або комплектація — перевитрата (плінтус, комплект септика, з'єднувач ринви)",
                  S1="Разом підтверджені складові",
                  grp2="Складові, що потребують підтвердження",
                  notused="Позиції кошторису, не списані на жоден будинок",
                  partial="Позиції, списані не на всі будинки (норма будинків без списання)",
                  incomp="Незіставні одиниці або комплектація — економія (гайка, ґрунт, шайба, обрізні диски)",
                  S2="Разом складові, що потребують підтвердження",
                  delta="Разом різниця з нормою (списано мінус норма)",
                  left="Ще лежить на складах цих будинків",
                  PROB="Імовірна економія — різниця з нормою разом із залишками на складах будинків (як перевитратою)",
                  LB="Нижня межа з запасом — підтверджені складові разом з усіма залишками на складах будинків (як перевитратою)")
w(f'<tr class="grp"><td colspan="4">{LBL["grp1"]}</td></tr>')
for key in ("price", "qty", "out", "comp"):
    w(f'<tr class="sub"><td>{LBL[key]}</td>{td(col["15"][key])}{td(col["30"][key])}{td(tot[key])}</tr>')
w(f'<tr class="total"><td>{LBL["S1"]}</td>{td(col["15"]["S1"])}{td(col["30"]["S1"])}{td(tot["S1"])}</tr>')
w(f'<tr class="grp"><td colspan="4">{LBL["grp2"]}</td></tr>')
for key in ("notused", "partial", "incomp"):
    w(f'<tr class="sub"><td>{LBL[key]}</td>{td(col["15"][key])}{td(col["30"][key])}{td(tot[key])}</tr>')
w(f'<tr class="total"><td>{LBL["S2"]}</td>{td(col["15"]["S2"])}{td(col["30"]["S2"])}{td(tot["S2"])}</tr>')
w(f'<tr class="total"><td>{LBL["delta"]}</td>{td(col["15"]["delta"])}{td(col["30"]["delta"])}{td(tot["delta"])}</tr>')
w(f'<tr><td>{LBL["left"]}</td>{td(col["15"]["left"], neutral=True)}{td(col["30"]["left"], neutral=True)}{td(tot["left"], neutral=True)}</tr>')
w(f'<tr class="total"><td>{LBL["PROB"]}</td>{td(col["15"]["PROB"])}{td(col["30"]["PROB"])}'
  f'<td class="num good">{n(tot["PROB"], True)} ({pc(-tot["PROB"], N8)})</td></tr>')
w(f'<tr class="key"><td>{LBL["LB"]}</td>'
  f'<td class="num good">{n(col["15"]["LB"], True)} ({pc(-col["15"]["LB"], col["15"]["norm"])})</td>'
  f'<td class="num good">{n(col["30"]["LB"], True)} ({pc(-col["30"]["LB"], col["30"]["norm"])})</td>'
  f'<td class="num good">{n(tot["LB"], True)} ({pc(-tot["LB"], N8)})</td></tr>')
w('</tbody></table>')
w('<p>Обидві оцінки вважають залишки на складах будинків перевитратою. Здебільшого це несписана норма тих самих будинків (генератор '
  '15 м №4, плити ЦСП і ОСБ 30 м №2), тож нижню межу пораховано з запасом.</p>')
w(f'<p>У рядку «не списані на жоден будинок» {n(AL_NU)} грн — позиції, списані під іншою загальною назвою (дозатори — як диспенсери на 15 м; '
  f'кріплення ринви — як держак ринви, кутики сталеві — як кутики (металопрокат) на 30 м); ці ж позиції дають {n(AL_OUT)} грн у рядку '
  f'«списано матеріалів, яких немає в кошторисі». З урахуванням цього нижня межа — {n(LB_AL)} грн ({pc(LB_AL, N8)}).</p>')

# --- Таблиця цін
w("<h3>Найбільша економія за ціною закупівлі — 15 м №1–№6</h3>")
w('<table><thead><tr><th>Загальна назва</th><th>Ціна: кошторис 15 м → фактична, грн/од.</th><th>Дешевше на, грн/од.</th>'
  '<th>Списано на 6 будинків, од. (% норми)</th><th>Економія за ціною, грн</th></tr></thead><tbody>')
for nmu, pss, pf, d, dp, cq, cqp, ek in price_rows:
    w(f'<tr><td>{html.escape(nmu)}</td><td class="num">{n(pss)} → {n(pf)}</td><td class="num">{n(d)} ({pct(dp, 0)})</td>'
      f'<td class="num">{n(cq)} ({pct(cqp, 0)})</td><td class="num good">{n(ek)}</td></tr>')
w(f'<tr class="total"><td colspan="4">Разом за {len(price_rows)} позиціями — {pc(PT_SUM, -col["15"]["price"], 0)} економії за ціною на 15 м '
  f'({n(-col["15"]["price"])} грн)</td><td class="num good">{n(PT_SUM)}</td></tr></tbody></table>')
w('<p>Дорожче за кошторис (15 м, за ціною): ' + "; ".join(f'{nm(r["name"])} — +{n(r["d15_pe"])}' for r in dear) + ' грн.</p>')
w(f'<p>Фактична ціна — за документами закупівлі матеріалів, списаних на ці будинки, з ПДВ. У листі від 30.09 — середня ціна кошторису '
  f'13 будинків і середня ціна всіх закупівель, тому цифри інші (генератор: там {n(gen["n13_s"]/gen["n13_q"])} → {n(GEN_PUR_P)}, '
  f'тут {n(ON[gen["k"]]["nr|15-1|ПланГрн"])} → {n(GEN15_P)}).</p>')
w(f'<p>За обсягом на всіх будинках стабільно списують менше за норму: профнастилу — {n(pn["d15_cq"]/6)} м² на будинок 15 м '
  f'(норма — {n(ON[pn["k"]]["nr|15-1|ПланКол"])} м²) і {n(pn["d30_cq"]/2)} м² на будинок 30 м (норма — {n(ON[pn["k"]]["nr|30-1|ПланКол"])} м²); '
  f'плит ЦСП на 15 м — {fq(cs["d15_cq"]/6)} шт. на будинок (норма — {fq(ON[cs["k"]]["nr|15-1|ПланКол"])} шт.).</p>')

# --- Чому 3,5 %
w("<h2>Чому щотижневий план-факт показує 3,5 %</h2>")
w(f'<p>Звіт «План-факт виконання СС» враховує закупівлі всього проєкту (13 будинків) і додає до них плановий розрахунок матеріалів — '
  f'незакуплений залишок кошторису за ціною СС; разом це «план на факт» звіту — {n(PNF)} грн при плані СС {n(PLAN13)} грн. Тому економію '
  f'відвантажених будинків у ньому перекриває решта прогнозу:</p>')
w('<table><thead><tr><th>Складова відхилення до кошторису у звіті</th><th>Грн (− економія, + перевитрата)</th></tr></thead><tbody>')
w(f'<tr><td>15 м №1–№6: фактичне списання проти норми</td>{td(col["15"]["delta"])}</tr>')
w(f'<tr><td>30 м №1–№2: фактичне списання проти норми</td>{td(col["30"]["delta"])}</tr>')
w(f'<tr><td>ще лежить на складах цих 8 будинків</td>{td(tot["left"], neutral=True)}</tr>')
w(f'<tr class="total"><td>Імовірна економія 8 відвантажених будинків</td>{td(tot["PROB"])}</tr>')
w('<tr class="grp"><td colspan="2">Решта прогнозу звіту: незавершені 30 м №3–№7 та інше</td></tr>')
w(f'<tr class="sub"><td>закуплене, крім списаного на 8 будинків і залишків на їхніх складах ({n(PUR)} − {n(tot["cons"])} − {n(tot["left"])})</td>{td(NOTB - tot["left"], neutral=True)}</tr>')
w(f'<tr class="sub"><td>плановий розрахунок матеріалів — незакуплений залишок кошторису за ціною СС</td>{td(ETC, neutral=True)}</tr>')
w(f'<tr class="sub"><td>норма кошторису СС 30 м №3–№7 (5 × 1\u00a0009\u00a0864,74)</td>{td(-N5, neutral=True)}</tr>')
w(f'<tr class="total"><td>Разом решта прогнозу звіту</td>{td(REST2, neutral=True)}</tr>')
w(f'<tr class="key"><td>Разом відхилення до кошторису у звіті на 30.09.2026 ({pc(DEV, PLAN13)})</td>{td(DEV)}</tr></tbody></table>')
w('<h3>Де зараз закуплені матеріали</h3>')
w('<table><thead><tr><th>Показник</th><th>Грн</th></tr></thead><tbody>')
w(f'<tr class="total"><td>Закуплено для проєкту (13 будинків)</td><td class="num">{n(PUR)}</td></tr>')
w(f'<tr class="sub"><td>списано на 8 відвантажених будинків</td><td class="num">{n(tot["cons"])}</td></tr>')
w(f'<tr class="sub"><td>ще лежить на складах цих 8 будинків</td><td class="num">{n(tot["left"])}</td></tr>')
w(f'<tr class="sub"><td>30 м №3: списано й лежить на складі будинку</td><td class="num">{n(H303)}</td></tr>')
w(f'<tr class="sub"><td><b>лежить на складі проєкту «ІБ _МД IRC 2026»</b></td><td class="num"><b>{n(STOCK_P)}</b></td></tr>')
w(f'<tr class="sub"><td>інше: сальдо інших складів — надійшло на них матеріалів IRC мінус отримано з них будинками й складом проєкту ({n(X_S)}); МШП IRC на загальному складі ({n(MSP_S, True)}); '
  f'додаткові витрати, віднесені на матеріали лише в бухобліку, за мінусом пального, списаного зі складу проєкту ({n(RESID)})</td><td class="num">{n(OTHER)}</td></tr>')
w(f'<tr><td>Плановий розрахунок матеріалів — незакуплений залишок кошторису за ціною СС</td><td class="num">{n(ETC)}</td></tr>')
w(f'<tr class="total"><td>Разом «план на факт» у звіті (закуплено + плановий розрахунок)</td><td class="num">{n(PNF)}</td></tr></tbody></table>')
w(f'<p>Решта прогнозу звіту (+{n(REST2)} грн) — не перевитрата на відвантажених будинках, а особливість розрахунку звіту й запас на складі проєкту. '
  f'Звіт вважає витраченим увесь матеріал на складі '
  f'проєкту й додає плановий розрахунок {n(ETC)} грн. Якщо порівняти норму СС 30 м №3–№7 із залишками складів за кожною позицією окремо '
  f'(надлишок одного матеріалу не покриває нестачі іншого), до норми бракує {n(DEF)} грн, тож плановий розрахунок більший за цю нестачу на '
  f'<b>{n(OVER)} грн</b> — це позиції, які на відвантажених будинках не знадобилися (бойлер 82, інформаційний банер, ґрунт), і матеріал, що вже '
  f'є на складі. Ще <b>{n(SURPLUS)} грн</b> матеріалів на складі проєкту — понад норму СС усіх незавершених будинків (з них {n(NONORM_P)} грн — '
  f'позиції, яких немає в кошторисі); якщо інвентаризація підтвердить цей запас і він не знадобиться, це надлишкова закупівля. '
  f'Натомість норма 30 м №3–№7 за цінами наявного запасу на {n(ECON37)} грн нижча за кошторисну; інше — {n(MISC)} грн (рядок «інше» '
  f'з таблиці, {n(OTHER)} грн, і залишок складу 30 м №3 понад його норму, {n(OWN_SURP, True)} грн). Разом: {n(OVER)} + {n(SURPLUS)} − {n(ECON37)} − {n(-MISC)} = {n(REST2)} грн.</p>')

# --- Чотири оцінки
w("<h2>Чотири оцінки економії — що показує кожна</h2>")
w('<table><thead><tr><th>Джерело</th><th style="text-align:left">Що показує</th><th>Будинків</th><th>Економія, грн (% до кошторису цих будинків)</th></tr></thead><tbody>')
w(f'<tr><td>Щотижневий «План-факт виконання СС»</td><td class="txt">закупівлі + плановий розрахунок матеріалів за ціною СС</td>'
  f'<td class="num">13</td><td class="num good">{n(-DEV)} ({pc(-DEV, PLAN13)})</td></tr>')
w(f'<tr><td>«Списання за нормами і понад норму (бухоблік)», колонка «Відхилення»</td><td class="txt">списання по будинках; ПДВ — за ставкою з картки номенклатури</td>'
  f'<td class="num">8</td><td class="num good">{n(-NR_DEV)} ({pc(-NR_DEV, N8)})</td></tr>')
w(f'<tr><td>Бюджет руху коштів (прийнятий бюджет жовтень–грудень 2026)</td><td class="txt">оплачено + план оплат на жовтень–листопад</td>'
  f'<td class="num">13</td><td class="num good">{n(EC_EK)} ({pc(EC_EK, D(str(cv["Q106"])))} до ліміту бюджету)</td></tr>')
w(f'<tr><td>Це дослідження</td><td class="txt">списання по будинках; ПДВ — фактичний, з документа закупівлі</td>'
  f'<td class="num">8</td><td class="num good">різниця з нормою {n(-tot["delta"])} ({pc(-tot["delta"], N8)}); імовірна економія {n(-tot["PROB"])} '
  f'({pc(-tot["PROB"], N8)}); нижня межа з запасом {n(-tot["LB"])} ({pc(-tot["LB"], N8)})</td></tr></tbody></table>')
w(f'<p><b>Звіт за нормами</b> бере ставку ПДВ з картки номенклатури, а не з документа закупівлі. У картках лінолеуму, внутрішніх дверей і плінтуса '
  f'стоїть «Без ПДВ», хоча їх куплено з ПДВ, — через це звіт занижує факт на {n(VAT_DN)} грн. Частину матеріалів, навпаки, куплено без ПДВ, '
  f'а в картці стоїть ставка 20\u00a0%, — через це звіт завищує факт на {n(VAT_UP)} грн. Сумарно звіт занижує факт на {n(VAT_NET)} грн. '
  f'Частину цього компенсує розбіжність у протилежний бік — +{n(SEMI)} грн до факту звіту, здебільшого через двічі врахований напівфабрикат '
  f'комплекту септика на 30 м №1 (разом {n(VAT_NET - SEMI)} грн; ще 0,01 грн — округлення у звіті). Колонка «Економія» того самого звіту '
  f'({n(NR_EK)} грн на 8 будинків) — інша міра: недовитрата за кількістю, оцінена за ціною кошторису; економії від дешевшої закупівлі в ній немає, '
  f'а перевитрати з неї не віднімаються. Її підсумок за групою «МД IRC 2026 30 м» — {n(NR_EK30)} грн — майже повністю припадає на незавершені '
  f'30 м №3–№7 ({n(NR_EK3_7)} грн), тобто це не економія.</p>')
w(f'<p><b>Бюджет руху коштів</b>: оплачено на 30.09 — {n(EC_PAID)} грн; на жовтень–листопад закладено {n(EC_BUD)} грн — {n(EC_DEBT)} грн боргів '
  f'і {n(EC_NEW)} грн нових закупівель; разом {n(EC_TOT)} грн. Від щотижневого звіту бюджет відрізняється тим, що на нові закупівлі закладено '
  f'на {n(NEW_D)} грн менше за плановий розрахунок ({n(ETC)} грн), а оплачене разом із боргами на {n(DEBT_X)} грн більше за закупівлі за обліком. '
  f'Локальний кошторис у книзі бюджету складено за тими самими рядками, що й кошторис СС, і майже з тими самими сумами: на 15 м — на {n(P15_LIM)} грн '
  f'більше, на 30 м — на {n(-P30_LIM)} грн менше, разом на {n(LIM_D)} грн менше. Звідси {n(-DEV)} + {n(NEW_D)} − {n(DEBT_X)} − {n(LIM_D)} = '
  f'{n(EC_EK)} грн: економія бюджету — переважно від урахування запасу на складі, а не від інших цін чи обсягів. Бюджет охоплює всі 13 будинків, '
  f'тобто й незавершені 30 м №3–№7.</p>')

# --- Що потрібно
w("<h2>Що потрібно, щоб підтвердити решту економії</h2>")
w('<ol>')
w(f'<li><b>Залишок на складі проєкту «ІБ _МД IRC 2026» — {n(STOCK_P)} грн за обліком.</b> Інвентаризація: чи є матеріал фізично й у тих самих '
  f'одиницях, що в обліку, — зокрема {n(SURPLUS)} грн понад норму СС незавершених будинків. Від неї залежить, скільки насправді треба докупити '
  f'для 30 м №3–№7: за обліковими залишками — {n(DEF)} грн, а не {n(ETC)} грн.</li>')
w(f'<li><b>Додаткове списання на відвантажені будинки.</b> На складах будинків лежить {n(tot["left"])} грн — здебільшого позиції, ще не списані на '
  f'ці будинки: генератор на складі 15 м №4 ({n(GEN4)} грн), плити ЦСП ({fq(CSP2[1])} шт., {n(CSP2[0])} грн) і ОСБ ({fq(OSB2[1])} шт., '
  f'{n(OSB2[0])} грн) на складі 30 м №2. Позицій, списаних не на всі будинки, — на {n(-tot["partial"])} грн за нормою; серед них також '
  f'обігрівачі й пінопласт 50 на 30 м №2, кондиціонер і гладкий лист на 30 м №1, а також комплект септика на 15 м №5 '
  f'(норма {n(-PND["15"])} грн — на будинок не списано жодного елемента септика). Для 30 м №1–№2 відмітку «Списання підтверджено» '
  f'в розрахунках комплектацій ще не поставлено.</li>')
w(f'<li><b>Генератор — найбільша економія за ціною: {n(GEN_PE)} грн.</b> У кошторисі «Генератор 7 кВт» — 70\u00a0000,00 грн (15 м) і '
  f'55\u00a0000,00 грн (30 м); закуплено {fq(gen["pu_q"])} шт. у середньому по {n(GEN_PUR_P)} грн. Потужність у документах закупівлі не вказана; '
  f'якщо вона нижча за 7 кВт, це інша комплектація, а не дешевша закупівля. Ці {n(GEN_PE)} грн входять у нижню межу; без генератора вона — '
  f'{n(GEN_LB)} грн ({pc(GEN_LB, N8)}).</li>')
w(f'<li><b>Позиції кошторису, не списані на жоден будинок, — {n(-tot["notused"])} грн;</b> найбільші: на 15 м — бойлер 82, інформаційний банер, ґрунт, '
  f'підвіконник ПВХ; на 30 м — стільці, бойлер 82, інформаційний банер. Якщо цих позицій у будинках немає, для цього типу будинку вони '
  f'в кошторисі зайві, і економія на цих позиціях підтверджується. Решта — у додатку, аркуш «По назвах», категорія «не списано на жоден будинок».</li>')
w(f'<li><b>Незіставні одиниці або комплектація — {n(INCOMP_TXT)} грн:</b> «Гайка» на 15 м ({n(GK15)} грн) — гайка М48 у кошторисі 15 м '
  f'по {n(m48["15"])} грн/шт, а в кошторисі 30 м — по {n(m48["30"])} грн/шт; «Шайба» ({n(SH)} грн) — шайба М12 у кошторисі по {n(sh12)} грн/шт; '
  f'ґрунт на 30 м ({n(GR30)} грн) — {fq(gr["d30_cq"])} кг списано лише на 30 м №1, по {n(gr["d30_cg"]/gr["d30_cq"])} грн/кг, а норма — '
  f'{fq(ON[gr["k"]]["nr|30-1|ПланКол"])} кг на будинок по {n(gr["d30_Ns"]/gr["d30_Nq"])} грн/кг; обрізні диски на 15 м ({n(DS15)} грн) — '
  f'{fq(ds["d15_cq"])} шт. списано лише на 15 м №1, по {n(ds["d15_cg"]/ds["d15_cq"])} грн, а норма — {fq(ON[ds["k"]]["nr|15-1|ПланКол"])} шт. '
  f'на будинок по {n(ds["d15_Ns"]/ds["d15_Nq"])} грн. Зіставність цих позицій із кошторисом документами не підтверджена.</li>')
w('</ol>')
assert INC_PARTS == INCOMP_TXT, (INC_PARTS, INCOMP_TXT)
w('<h3>Розбіжності в даних обліку</h3><ul>')
w(f'<li>Плінтус 15 м: у кошторисі СС після перерахунку одиниць — {fq(PL_TQ)} пог.\u00a0м на будинок по {n(PL_TS/PL_TQ)} грн, в оригіналі кошторису — '
  f'{fq(PL_OQ)} пог.\u00a0м по {n(PL_OS/PL_OQ)} грн (сума та сама). Через це щотижневий звіт показує перевитрату плінтуса за обсягом.</li>')
w('<li>Реалізацію ІБ00-000178 від 01.10.2026 вказано в розрахунках комплектацій і 30 м №2, і 30 м №3: відвантаження двох будинків '
  'прив\'язано до одного документа.</li></ul>')

# --- Як читати додаток
xl_name = "Додаток_економія_матеріалів_МД_IRC_2026_02-10-2026.xlsx"
w("<h2>Як читати додаток</h2>")
w(f'<div class="guide"><p>У файлі «{xl_name}» п\'ять аркушів. На перших трьох мінус — економія, плюс — перевитрата.</p><ul>'
  '<li><b>«З чого складається економія»</b> — та сама таблиця, що й у записці: 15 м, 30 м №1–№2, разом; в останній колонці — звідки ці суми '
  '(аркуш «По назвах», для залишків — аркуш «Будинки»).</li>'
  f'<li><b>«Будинки»</b> — для кожного будинку норма СС, скільки списано, різниця й залишок на складі будинку. Дві колонки «Списано»: перша — '
  f'як у звіті за нормами (ПДВ за карткою номенклатури), друга — з ПДВ за документом закупівлі; різницю з нормою обчислено за другою. На 15 м '
  f'списано від {n(min(cons[h] for h in H15))} до {n(max(cons[h] for h in H15))} грн на будинок (найменше — №4, на ньому не списано генератор).</li>'
  '<li><b>«По назвах»</b> — один рядок = одна загальна назва. Читається зліва направо, окремо для 15 м (6 будинків) і 30 м (2 будинки): '
  'норма на будинок → скільки списано в середньому на будинок → ціна за кошторисом → фактична ціна → різниця в гривнях і з чого вона '
  'складається: «у т. ч. за ціною» (купили дешевше чи дорожче) і «у т. ч. за обсягом» (списали менше чи більше); колонка '
  '«у т. ч. норма будинків без списання» показує, яка частина обсягу — норма будинків, на які позицію не списано. Колонки «Категорія 15 м» '
  'і «Категорія 30 м» показують, чи підтверджена цифра документами: «зіставно» — так; «списано не на всі будинки» — ціна й обсяг на будинках, '
  'де позицію списано, підтверджені, а норма будинків без списання потребує підтвердження; «не списано на жоден будинок», «незіставні одиниці» — '
  'потребують підтвердження; «інша одиниця / комплектація» (плінтус, комплект септика, з\'єднувач ринви) — перевитрата, врахована як підтверджена; '
  '«поза кошторисом» — списано, але в кошторисі немає; «—» — для цього типу будинку позиції немає ні в кошторисі, ні в списанні. Під рядком '
  '«Разом» — рядок «у т. ч. не розкладено на ціну й обсяг»: у сумі з колонками «у т. ч. за ціною» і «у т. ч. за обсягом» він дає повну різницю.</li>'
  '<li><b>«Куди пішли матеріали»</b> — закуплене для проєкту: скільки списано, скільки лежить на складах будинків і на складі проєкту, '
  'а внизу — скільки бракує до норми на 30 м №3–№7 за обліковими залишками і на скільки плановий розрахунок більший.</li>'
  '<li><b>«Порівняння економії»</b> — чотири джерела оцінки; для цього дослідження три рядки: різниця з нормою, імовірна економія, нижня межа з запасом; '
  'тут економія — додатне число.</li></ul>'
  f'<p>Приклад рядка «Профнастил покрівельний», 15 м: за кошторисом {n(exo["nr|15-1|ПланКол"])} м² на будинок по {n(exo["nr|15-1|ПланГрн"]/exo["nr|15-1|ПланКол"])} грн, '
  f'списано {n(ex["d15_cq"]/6)} м² на будинок по {n(ex["d15_cg"]/ex["d15_cq"])} грн. Різниця на 6 будинках {n(ex["d15_delta"], True)} грн: '
  f'за ціною {n(ex["d15_pe"], True)}, за обсягом {n(ex["d15_qe"], True)} — тобто економія переважно за обсягом.</p></div>')

w(f'<p class="src">Джерела: норма — картки структури собівартості (ERP і бухоблік збігаються рядок у рядок); закупівлі — ERP, як у щотижневому звіті '
  f'«План-факт виконання СС» на 30.09.2026 (його детальний варіант в 1С — «План на факт за матеріалами»: колонка «Факт, грн» — {n(PUR)} грн, '
  f'збігається до копійки; «Плановий розрахунок матеріалів, грн.» — {n(ETC)} грн; разом — «План на факт, грн» {n(PNF)} грн); списання й залишки — '
  f'бухоблік, ПДВ — фактичний, з документа закупівлі партії; звіт за нормами — «Списання IRC.xlsx» за 01.08–31.10.2026, сума трьох рядків колонки '
  f'«Відхилення»; бюджет — «Бюджет_Жовтень26-Грудень26.xlsx», аркуш «МД IRC 2026», рядок «Строительные материалы». Цифри звітів і бюджету '
  f'наведено як у файлах; решта — розрахунок за даними обліку.</p>')
w("</div>\n</body>\n</html>\n")
HTML = "\n".join(H)
# нерозривні пробіли між числом і одиницею в тексті (поза тегами)
_b = HTML.index("<body>")
_head, _body = HTML[:_b], HTML[_b:]
_body = re.sub(r"(?<=[\d>]) (грн|м²|м³|кг|шт\.|кВт|пог\.|%)", "\u00a0\\1", _body)
_body = re.sub(r"\b(15|30) м\b", "\\1\u00a0м", _body)
_body = _body.replace("у т. ч.", "у\u00a0т.\u00a0ч.")
HTML = _head + _body


# ---------- ДОДАТОК XLSX (укр.)
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
wb = openpyxl.Workbook(); wb.remove(wb.active)
BF = Font(bold=True); HF = PatternFill("solid", fgColor="DDEBF7"); TF = PatternFill("solid", fgColor="F2F2F2"); GF = PatternFill("solid", fgColor="EEF3F8")
GOOD = Font(color="1B7F4B"); BAD = Font(color="B3261E")

def sheet(title, head, rows, widths, note=None, money_cols=(), pct_cols=(), signed_cols=(), bold_rows=(), grp_rows=()):
    ws = wb.create_sheet(title); r0 = 1
    if note:
        c = ws.cell(1, 1, note); c.font = Font(italic=True, color="555555"); r0 = 3
    for j, h_ in enumerate(head, 1):
        c = ws.cell(r0, j, h_); c.font = BF; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, Decimal): v = float(v)
            c = ws.cell(i, j, v)
            if isinstance(v, float):
                c.number_format = "0.0" if j in pct_cols else ("#,##0.00" if j in money_cols else ("#,##0" if float(v).is_integer() else "#,##0.###"))
                if j in signed_cols:
                    if v < 0: c.font = GOOD
                    elif v > 0: c.font = BAD
        k = i - r0 - 1
        if k in bold_rows or k in grp_rows:
            for j in range(1, len(head) + 1):
                ws.cell(i, j).fill = GF if k in grp_rows else TF
                ws.cell(i, j).font = Font(bold=True, color=ws.cell(i, j).font.color)
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    for j, wd in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = wd
    ws.row_dimensions[r0].height = 48
    return ws

# 1. З чого складається економія
CAT = {"сопоставимо": "зіставно", "не использовано": "не списано на жоден будинок", "поза кошторисом": "поза кошторисом", "—": "—"}
REF = OrderedDict(price="«зіставно», «списано не на всі будинки» → «у т. ч. за ціною»",
                  qty="«зіставно», «списано не на всі будинки» → «у т. ч. за обсягом» мінус «у т. ч. норма будинків без списання»",
                  out="«поза кошторисом» → «різниця»",
                  comp="«інша одиниця / комплектація» → «різниця» мінус «у т. ч. норма будинків без списання» (септик 15 м)",
                  notused="«не списано на жоден будинок» → «різниця»",
                  partial="«списано не на всі будинки» і септик 15 м → «у т. ч. норма будинків без списання»",
                  incomp="«незіставні одиниці» → «різниця»",
                  delta="рядок «Разом»: колонки «15 м: різниця на 6 будинках», «30 м: різниця на 2 будинках», «Разом різниця, 8 будинків»",
                  left="аркуш «Будинки», колонка «Залишок на складі будинку»")
rows_z = [["Норма кошторису СС", col["15"]["norm"], col["30"]["norm"], N8, ""],
          ["Списано фактично (з ПДВ за документом закупівлі)", col["15"]["cons"], col["30"]["cons"], tot["cons"], ""],
          [LBL["grp1"], None, None, None, ""]]
for key in ("price", "qty", "out", "comp"):
    rows_z.append(["  " + LBL[key], col["15"][key], col["30"][key], tot[key], REF[key]])
rows_z.append([LBL["S1"], col["15"]["S1"], col["30"]["S1"], tot["S1"], ""])
rows_z.append([LBL["grp2"], None, None, None, ""])
for key in ("notused", "partial", "incomp"):
    rows_z.append(["  " + LBL[key], col["15"][key], col["30"][key], tot[key], REF[key]])
rows_z.append([LBL["S2"], col["15"]["S2"], col["30"]["S2"], tot["S2"], ""])
rows_z.append([LBL["delta"], col["15"]["delta"], col["30"]["delta"], tot["delta"], REF["delta"]])
rows_z.append([LBL["left"], col["15"]["left"], col["30"]["left"], tot["left"], REF["left"]])
rows_z.append([LBL["PROB"], col["15"]["PROB"], col["30"]["PROB"], tot["PROB"], ""])
rows_z.append([LBL["LB"], col["15"]["LB"], col["30"]["LB"], tot["LB"], ""])
sheet("З чого складається економія", ["Складова", "15 м №1–№6, грн", "30 м №1–№2, грн", "Разом, грн", "Звідки суми (аркуш «По назвах»: категорія → колонка; для залишків — аркуш «Будинки»)"],
      rows_z, [78, 16, 16, 16, 70], "Мінус — економія, плюс — перевитрата. Та сама таблиця, що в аналітичній записці.",
      money_cols=(2, 3, 4), signed_cols=(2, 3, 4), bold_rows=(7, 12, 13, 15, 16), grp_rows=(2, 8))

# 2. Будинки
rkm = {}
for x in RK:
    t = "15" if "15 м" in x["Спец"] else "30"
    m_ = re.search(r"№\s*(\d+)", x["Подр"]); rkm[f"{t}-{m_.group(1)}"] = x
fd = lambda s_: (lambda d_: f"{d_[8:10]}.{d_[5:7]}.{d_[0:4]}")(str(s_)[:10]) if s_ else ""
rows_b = []
for h in H15 + H30:
    d = F["house"][h]; x = rkm.get(h, {})
    cg = q2(cons[h]); nrm = q2(d["ПланГрн"])
    rows_b.append([h.replace("-", " м №"), nrm, NR_O[h], cg, cg - nrm, float(dv(cg - nrm, nrm) * 100), q2(left[h]), fd(x.get("D")),
                   (x.get("Реал", "") or "").split(" от ")[0].replace("Реалізація товарів і послуг ", ""), "так" if x.get("Подтв") else "ні"])
tb = ["Разом 8 будинків", sum(r[1] for r in rows_b), sum(r[2] for r in rows_b), sum(r[3] for r in rows_b), sum(r[4] for r in rows_b),
      float(dv(sum(r[4] for r in rows_b), sum(r[1] for r in rows_b)) * 100), sum(r[6] for r in rows_b), "", "", ""]
assert tb[3] == tot["cons"] and tb[1] == N8 and tb[6] == tot["left"] and tb[2] == NR_FACT and tb[4] == tot["delta"]
sheet("Будинки", ["Будинок", "Норма СС, грн", "Списано, грн — як у звіті за нормами (ПДВ за карткою номенклатури)", "Списано, грн — з ПДВ за документом закупівлі",
                  "Різниця з нормою, грн (за другою колонкою «Списано»)", "Різниця, %", "Залишок на складі будинку, грн", "Дата відвантаження",
                  "Реалізація", "Списання підтверджено в розрахунку комплектації"],
      rows_b + [tb], [16, 15, 22, 20, 20, 10, 16, 13, 13, 18],
      "Мінус — економія. 30 м №1–№2 відвантажені замовнику, акти виставлені; 30 м №3–№7 не завершені й не аналізуються.",
      money_cols=(2, 3, 4, 5, 7), pct_cols=(6,), signed_cols=(5, 6), bold_rows=(8,))

# 3. По назвах
def lr(vals, target):
    rd = [q2(v) for v in vals]; diff = int((target - sum(rd)) * 100)
    order = sorted(range(len(vals)), key=lambda i: (vals[i] - rd[i]) * (1 if diff > 0 else -1), reverse=True)
    for i in order[:abs(diff)]: rd[i] += Decimal("0.01") * (1 if diff > 0 else -1)
    if diff: LOG.append(f"найбільший залишок: {diff} коп.")
    return rd
items = [r for r in sorted(R.values(), key=lambda r: r["d15_delta"] + r["d30_delta"]) if r["d15_delta"] or r["d30_delta"]]
cat = {(r["k"], g): classify(r, g) for r in items for g in ("d15", "d30")}
alloc_ = {}
for g, hs in (("d15", H15), ("d30", H30)):
    gg = g[1:]
    grp_t = {"сопоставимо": col[gg]["price"] + col[gg]["qty"] + col[gg]["partial"] - PND[gg], "поза кошторисом": col[gg]["out"],
             "не использовано": col[gg]["notused"], "nd": col[gg]["comp"] + col[gg]["incomp"] + PND[gg]}
    dl = [Z] * len(items)
    for key_, tgt in grp_t.items():
        idx = [i for i, r in enumerate(items) if (cat[(r["k"], g)] in ("единица/состав", "несопоставимо") if key_ == "nd" else cat[(r["k"], g)] == key_)]
        for i, v in zip(idx, lr([items[i][f"{g}_delta"] for i in idx], tgt)): dl[i] = v
    assert sum(dl) == col[gg]["delta"]
    comp_i = [i for i, r in enumerate(items) if cat[(r["k"], g)] == "сопоставимо"]
    pe = lr([items[i][f"{g}_pe"] for i in comp_i], col[gg]["price"])
    pzx = []
    for i in comp_i:
        o = ON[items[i]["k"]]; z = sum(1 for h in hs if o[f"c|{h}|q"] == 0)
        pzx.append(-(Decimal(z) * o[f"nr|{hs[0]}|ПланГрн"]) if 0 < z < len(hs) else Z)
    pz = lr(pzx, col[gg]["partial"] - PND[gg])
    pe_full = [None] * len(items); qe_full = [None] * len(items); pz_full = [None] * len(items)
    for i, v, zv in zip(comp_i, pe, pz):
        pe_full[i] = v; qe_full[i] = dl[i] - v; pz_full[i] = zv if zv else None
    for i, r in enumerate(items):
        if cat[(r["k"], g)] in ("поза кошторисом", "не использовано"): pe_full[i] = Z; qe_full[i] = dl[i]
        if cat[(r["k"], g)] in ("единица/состав", "несопоставимо"):
            o = ON[r["k"]]; z = sum(1 for h in hs if o[f"c|{h}|q"] == 0)
            pzv = (Decimal(z) * o[f"nr|{hs[0]}|ПланГрн"]) if 0 < z < len(hs) else Z
            if pzv and r[f"{g}_delta"] + pzv > 0: pz_full[i] = -q2(pzv)
    alloc_[g] = (dl, pe_full, qe_full, pz_full)
rows_n = []
for i, r in enumerate(items):
    o = ON[r["k"]]
    def blk(g, h0, nh):
        nq_ = o[f"nr|{h0}|ПланКол"]; ns = o[f"nr|{h0}|ПланГрн"]; dl, pe_, qe_, pz_ = alloc_[g]
        return [nq_ if nq_ else None, (r[f"{g}_cq"] / nh) if r[f"{g}_cq"] else None, dv(ns, nq_) if nq_ else None,
                dv(r[f"{g}_cg"], r[f"{g}_cq"]) if r[f"{g}_cq"] else None, dl[i], pe_[i], qe_[i], pz_[i]]
    def lab(g, nh):
        c = cat[(r["k"], g)]
        if c == "сопоставимо":
            z = sum(1 for h in (H15 if g == "d15" else H30) if o[f"c|{h}|q"] == 0)
            return "списано не на всі будинки" if 0 < z < nh else "зіставно"
        if c in ("единица/состав", "несопоставимо"):
            return "інша одиниця / комплектація" if r[f"{g}_delta"] > 0 else "незіставні одиниці"
        if not r[f"{g}_delta"]: return "—"
        return CAT[c]
    b15, b30 = blk("d15", "15-1", 6), blk("d30", "30-1", 2)
    rows_n.append([nm(r["name"]), unit(r)] + b15 + b30 + [b15[4] + b30[4], lab("d15", 6), lab("d30", 2)])
# колонки: 1 назва, 2 од., 3–10 = 15 м (норма, сер., ціна СС, ціна факт, різниця, ціна, обсяг, норма без списання), 11–18 = 30 м, 19 разом, 20–21 категорії
S_ = lambda j: sum((x[j] for x in rows_n if isinstance(x[j], Decimal)), Z)
tn = ["Разом", ""] + [""] * 4 + [S_(6), S_(7), S_(8), S_(9)] + [""] * 4 + [S_(14), S_(15), S_(16), S_(17), S_(18), "", ""]
nd = ["у т. ч. не розкладено на ціну й обсяг (незіставні одиниці, інша одиниця / комплектація)", ""] + [""] * 4 + \
     [S_(6) - S_(7) - S_(8), "", "", ""] + [""] * 4 + [S_(14) - S_(15) - S_(16), "", "", "", S_(18) - S_(7) - S_(8) - S_(15) - S_(16), "", ""]
assert tn[18] == tot["delta"] and tn[6] == col["15"]["delta"] and tn[14] == col["30"]["delta"]
assert tn[6] - tn[7] - tn[8] == col["15"]["comp"] + col["15"]["incomp"] + PND["15"], (tn[6] - tn[7] - tn[8], col["15"]["comp"] + col["15"]["incomp"] + PND["15"])
assert tn[14] - tn[15] - tn[16] == col["30"]["comp"] + col["30"]["incomp"] + PND["30"]
assert tn[9] == col["15"]["partial"] and tn[17] == col["30"]["partial"]
H_N = ["Загальна назва", "Од.",
       "15 м: норма на будинок", "15 м: списано на будинок (середнє)", "15 м: ціна за кошторисом, грн", "15 м: фактична ціна, грн",
       "15 м: різниця на 6 будинках, грн", "у т. ч. за ціною", "у т. ч. за обсягом", "у т. ч. норма будинків без списання",
       "30 м: норма на будинок", "30 м: списано на будинок (середнє)", "30 м: ціна за кошторисом, грн", "30 м: фактична ціна, грн",
       "30 м: різниця на 2 будинках, грн", "у т. ч. за ціною", "у т. ч. за обсягом", "у т. ч. норма будинків без списання",
       "Разом різниця, 8 будинків, грн", "Категорія 15 м", "Категорія 30 м"]
sheet("По назвах", H_N, rows_n + [tn, nd], [34, 7, 11, 11, 11, 11, 14, 13, 13, 13, 11, 11, 11, 11, 14, 13, 13, 13, 15, 26, 26],
      "Мінус — економія, плюс — перевитрата. Ціни — з ПДВ. «За ціною» = списана кількість × (фактична ціна − ціна за кошторисом); "
      "«за обсягом» = (списано − норма) × ціна за кошторисом; «норма будинків без списання» — частина «за обсягом» для будинків, на які позицію "
      "не списано. Для категорій «не списано на жоден будинок» і «поза кошторисом» уся різниця — в обсязі; для категорій «незіставні одиниці» "
      "та «інша одиниця / комплектація» різниця на ціну й обсяг не розкладається (див. рядок під «Разом»).",
      money_cols=(5, 6, 7, 8, 9, 10, 13, 14, 15, 16, 17, 18, 19), signed_cols=(7, 8, 9, 10, 15, 16, 17, 18, 19),
      bold_rows=(len(rows_n), len(rows_n) + 1))

# 4. Куди пішли матеріали
rows_k = [["Закуплено для проєкту (13 будинків), ERP", PUR],
          ["  списано на 8 відвантажених будинків", tot["cons"]],
          ["  ще лежить на складах цих 8 будинків", tot["left"]],
          ["  30 м №3: списано й лежить на складі будинку", H303],
          ["  лежить на складі проєкту «ІБ _МД IRC 2026»", STOCK_P],
          [f"  інше: сальдо інших складів — надійшло на них матеріалів IRC мінус отримано з них будинками й складом проєкту ({n(X_S)}); МШП IRC на загальному складі ({n(MSP_S, True)}); "
           f"додаткові витрати, віднесені на матеріали лише в бухобліку, за мінусом пального, списаного зі складу проєкту ({n(RESID)})", OTHER],
          ["Плановий розрахунок матеріалів — незакуплений залишок кошторису за ціною СС", ETC],
          ["Разом «план на факт» у щотижневому звіті (закуплено + плановий розрахунок)", PNF],
          ["Скільки бракує до норми на 30 м №3–№7 (за обліковими залишками складів, за кожною позицією)", None],
          ["  30 м №3", DEF3], ["  30 м №4–№7", DEF47], ["Разом бракує до норми", DEF],
          ["Плановий розрахунок більший за нестачу за обліковими залишками на", OVER],
          ["Склад проєкту понад норму СС усіх незавершених будинків", SURPLUS],
          ["  у т. ч. позиції, яких немає в кошторисі", NONORM_P]]
sheet("Куди пішли матеріали", ["Показник", "Грн, з ПДВ"], rows_k, [104, 18],
      "Закупівлі — ERP по 30.09.2026; списання й залишки — бухоблік на 02.10.2026.", money_cols=(2,), bold_rows=(0, 7, 11, 12, 13), grp_rows=(8,))

# 5. Порівняння економії
rows_t = [["Щотижневий «План-факт виконання СС»", "закупівлі + плановий розрахунок матеріалів за ціною СС", 13, -DEV, float(dv(-DEV, PLAN13) * 100)],
          ["«Списання за нормами і понад норму (бухоблік)», колонка «Відхилення»", "списання по будинках; ПДВ — за ставкою з картки номенклатури", 8, -NR_DEV, float(dv(-NR_DEV, N8) * 100)],
          ["Бюджет руху коштів (прийнятий бюджет жовтень–грудень 2026)", "оплачено + план оплат на жовтень–листопад; % — до ліміту бюджету", 13, EC_EK, float(dv(EC_EK, D(str(cv["Q106"]))) * 100)],
          ["Це дослідження: різниця з нормою", "списання по будинках; ПДВ — фактичний, з документа закупівлі", 8, -tot["delta"], float(dv(-tot["delta"], N8) * 100)],
          ["Це дослідження: імовірна економія", "різниця з нормою мінус залишки на складах будинків", 8, -tot["PROB"], float(dv(-tot["PROB"], N8) * 100)],
          ["Це дослідження: нижня межа з запасом", "без позицій, що потребують підтвердження; залишки на складах будинків віднято повністю як перевитрату", 8, -tot["LB"], float(dv(-tot["LB"], N8) * 100)]]
sheet("Порівняння економії", ["Джерело", "Що показує", "Будинків", "Економія, грн", "Економія, % до кошторису цих будинків"], rows_t, [62, 76, 10, 16, 14],
      "На цьому аркуші економія — додатне число.", money_cols=(4,), pct_cols=(5,))


# ---------- запис
WEEK = os.path.join(B, r"30092026\IRC 2026")
fn_h = "Аналітична_записка_МД_IRC_2026_економія_матеріалів_02-10-2026.html"
spec = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), r"..\..\..\notebook\specs"))
open(os.path.join(WEEK, fn_h), "w", encoding="utf-8").write(HTML)
wb.save(os.path.join(WEEK, xl_name))
open(os.path.join(spec, "2026-10-02-ekonomiya-materialiv-ss-irc-ZAPYSKA.html"), "w", encoding="utf-8").write(HTML)
wb.save(os.path.join(spec, "2026-10-02-ekonomiya-materialiv-ss-irc-DODATOK.xlsx"))
print("HTML:", os.path.join(WEEK, fn_h)); print("XLSX:", os.path.join(WEEK, xl_name))
for l in LOG: print("  ", l)
KEYS = dict(cons=tot["cons"], delta=tot["delta"], LB=tot["LB"], LB15=col["15"]["LB"], LB30=col["30"]["LB"], PROB=tot["PROB"], S1=tot["S1"], S2=tot["S2"],
            left=tot["left"], price=tot["price"], price15=col["15"]["price"], price30=col["30"]["price"], qty=tot["qty"], qty15=col["15"]["qty"],
            qty30=col["30"]["qty"], out=tot["out"], comp=tot["comp"], comp15=col["15"]["comp"], comp30=col["30"]["comp"], notused=tot["notused"],
            partial=tot["partial"], partial15=col["15"]["partial"], partial30=col["30"]["partial"], incomp=tot["incomp"], incomp15=col["15"]["incomp"],
            incomp30=col["30"]["incomp"], AL_NU=AL_NU, AL_OUT=AL_OUT, LB_AL=LB_AL, GEN_LB=GEN_LB, REST=REST2, SURPLUS=SURPLUS, NONORM_P=NONORM_P, ECON37=ECON37, MISC=MISC, NOTB=NOTB, OTHER=OTHER, DEF=DEF, OVER=OVER, GEN_PE=GEN_PE, PT=PT_SUM, EC_NEW=EC_NEW, EC_DEBT=EC_DEBT,
            NEW_D=NEW_D, DEBT_X=DEBT_X)
json.dump({k: str(v) for k, v in KEYS.items()}, open(os.path.join(OUT, "zapiska_keys.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("Ключові:", {k: n(v) for k, v in KEYS.items()})
