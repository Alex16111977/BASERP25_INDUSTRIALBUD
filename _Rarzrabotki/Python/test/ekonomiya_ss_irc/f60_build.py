# -*- coding: utf-8 -*-
"""Ф1–Ф5. Итоговые расчёты и ANALIZ.xlsx: таблица по загальним назвам (ф.1), ТОП-80 %, мост (ф.2), дома,
сверка трёх «экономий» (ф.3), деньги и смета экономиста (ф.4), корзины вывода (ф.5).
Источник — out/*.json (выгрузки f10…f53). Числа отчётов приводятся как есть; всё остальное — «расчёт»."""
import os, sys, json
sys.path.insert(0, os.path.dirname(__file__))
from common import *
from collections import defaultdict, OrderedDict

L = lambda n: json.load(open(os.path.join(OUT, n), encoding="utf-8"))
A = L("analysis.json")["rows"]; C = L("core.json"); MVP = L("money_vs_purch.json"); EV = L("economist_vs_ss.json")
RK = L("rk_irc.json"); NRR = [r for r in L("norm_report_rows.json")["rows"] if "IR" in (r["Подразделение"] or "")]
Z = Decimal(0)
R = {k: {f: (D(v) if isinstance(v, str) and _isnum(v) else v) for f, v in r.items()} for k, r in A.items()} if False else None

def isnum(s):
    try:
        Decimal(s); return True
    except Exception:
        return False

R = {k: {f: (D(v) if isinstance(v, str) and isnum(v) and f not in ("k", "name") else v) for f, v in r.items()} for k, r in A.items()}
ON = {k: defaultdict(Decimal, {m: D(v) for m, v in d.items()}) for k, d in C["on"].items()}
H15 = [f"15-{i}" for i in range(1, 7)]
H30 = ["30-1", "30-2"]
SC8 = H15 + H30          # границы по указанию заказчика 02.10.2026: 15 м №1–№6 + 30 м №1–№2 (отгружены, акты)

def dv(a, b):
    return a / b if b else Z

# ---------- экономист: цена/объём на дом по ОН
econ = defaultdict(lambda: defaultdict(Decimal))
for sh, t in (("IRS 15", "15"), ("IRS 30", "30")):
    for x in EV[sh]["rows"]:
        k = " ".join(str(x["ОН"] or "").split()).casefold()
        econ[k][f"q{t}"] += D(x["q"]); econ[k][f"s{t}"] += D(x["s"])

# ---------- классификация расхода 15 м / 30 м (ф.1, Б) с учётом «единица/состав»
def classify(r, g):
    cat = r[f"{g}_cat"]
    Nq, Cq = r[f"{g}_Nq"], r[f"{g}_cq"]
    Ns, Cg = r[f"{g}_Ns"], r[f"{g}_cg"]
    if cat == "сопоставимо":
        pr = dv(dv(Cg, Cq), dv(Ns, Nq)); qr = dv(Cq, Nq)
        if (pr < D("0.5") and qr > 2) or (pr > 2 and qr < D("0.5")):
            return "единица/состав"
    return cat

T = defaultdict(Decimal)
rows1 = []
for k, r in R.items():
    o = ON[k]
    c15 = classify(r, "d15"); c30 = classify(r, "d30")
    r["c15"] = c15; r["c30"] = c30
    pz = r["d15_partial_zero"] if c15 == "сопоставимо" else Z
    # нижняя граница по 15 м: экономии только надёжные; перерасходы — все
    if c15 == "сопоставимо":
        lb = r["d15_pe"] + r["d15_qe"] + pz
    elif c15 == "поза кошторисом":
        lb = r["d15_delta"]
    elif c15 in ("единица/состав", "несопоставимо"):
        lb = max(Z, r["d15_delta"])
    else:
        lb = Z
    r["lb15"] = lb
    nq30, ns30 = o["nr|30-1|ПланКол"], o["nr|30-1|ПланГрн"]
    z30 = sum(1 for h in H30 if o[f"c|{h}|q"] == 0)
    pz30 = (Decimal(z30) * nq30 * dv(ns30, nq30)) if (c30 == "сопоставимо" and 0 < z30 < 2) else Z
    r["d30_partial_zero"] = pz30; r["d30_zero_houses"] = z30
    if c30 == "сопоставимо":
        lb3 = r["d30_pe"] + r["d30_qe"] + pz30
    elif c30 == "поза кошторисом":
        lb3 = r["d30_delta"]
    elif c30 in ("единица/состав", "несопоставимо"):
        lb3 = max(Z, r["d30_delta"])
    else:
        lb3 = Z
    r["lb30"] = lb3
    for f in ("d15_delta", "d15_pe", "d15_qe", "d30_delta", "d30_pe", "d30_qe", "lb15", "lb30", "d30_partial_zero"):
        T[f] += r[f]
    T[f"d15|{c15}"] += r["d15_delta"]; T[f"d30|{c30}"] += r["d30_delta"]
    if c15 == "сопоставимо":
        T["d15_pe_rel"] += r["d15_pe"]; T["d15_qe_rel"] += r["d15_qe"]; T["d15_pz"] += pz
    if c30 == "сопоставимо":
        T["d30_pe_rel"] += r["d30_pe"]; T["d30_qe_rel"] += r["d30_qe"]
    r["etc_minus_def"] = r["etc"] - r["e303_def"] - r["e47_def"]
    e = econ.get(k, {})
    r["econ_p15"] = dv(e.get("s15", Z), e.get("q15", Z)); r["econ_q15"] = e.get("q15", Z)
    r["econ_p30"] = dv(e.get("s30", Z), e.get("q30", Z)); r["econ_q30"] = e.get("q30", Z)
    rows1.append(r)

# ---------- остатки 15 м и итоги домов из отчёта по нормам (как есть)
house = OrderedDict()
for r in NRR:
    t = "15" if "15 м" in r["Подразделение"] else "30"
    h = f"{t}-" + (r["Домик"] or "").replace("№", "").strip()
    d = house.setdefault(h, defaultdict(Decimal))
    if (r["ВидДокумента"] or "").startswith("6."):
        continue
    for c in ("ПланГрн", "ФактВНормеГрн", "ФактПонадГрн", "ЭкономияГрн", "ОстатокОбъектаГрн", "ОтклонениеГрн"):
        d[c] += D(r[c] or 0)
cons_gt = {h: sum(ON[k][f"c|{h}|gt"] for k in ON) for h in house}
left_gt = {h: sum(ON[k][f"stkh|{h}|gt"] for k in ON) for h in house}

# ---------- корзины
left15 = sum(left_gt[h] for h in H15)
left3012 = left_gt["30-1"] + left_gt["30-2"]
LB15 = -(T["lb15"]) - left15                     # экономия = −Δ
LB30 = -(T["lb30"]) - left3012
LB = LB15 + LB30
d15 = -T["d15_delta"]; d30 = -T["d30_delta"]
e303 = sum(r["e303_cost"] for r in rows1)
N15 = D("704612.01"); N30 = D("1009864.74"); N9 = N15 * 6 + N30 * 2; N13 = N15 * 6 + N30 * 7   # N9 = норма 8 домов в границах
probable = d15 - left15 + d30 - left3012
deficit = sum(r["e303_def"] + r["e47_def"] for r in rows1)
ETC = D("1849175.41"); PUR = D("9053180.08")
eac_fact13 = PUR + deficit
print(f"15м×6: норма {fmt(N15*6)} расход реальн.брутто {fmt(sum(cons_gt[h] for h in H15))} экономия {fmt(d15)} ({fmt(d15/(N15*6)*100,1)} %)")
print(f"   цена (надёжн.) {fmt(-T['d15_pe_rel'])} кол-во (надёжн.) {fmt(-T['d15_qe_rel'])} из них частичные нули {fmt(T['d15_pz'])}")
for c in ("сопоставимо", "единица/состав", "несопоставимо", "не использовано", "поза кошторисом", "—"):
    print(f"   категория {c:<16} Δ {fmt(T['d15|'+c])}")
print(f"   остатки на складах 15 м (реальн.брутто) {fmt(left15)}")
print(f"НИЖНЯЯ ГРАНИЦА: 15 м {fmt(LB15)} ({fmt(LB15/(N15*6)*100,1)} %), 30 м №1–2 {fmt(LB30)} ({fmt(LB30/(N30*2)*100,1)} %); ИТОГО {fmt(LB)} = {fmt(LB/N9*100,1)} % нормы 8 домов; {fmt(LB/N13*100,1)} % СС 13 домов")
print(f"   30м: цена (надёжн.) {fmt(-T['d30_pe_rel'])} кол-во (надёжн.) {fmt(-T['d30_qe_rel'])} частичные нули {fmt(T['d30_partial_zero'])}")
print(f"30м №1–2: экономия {fmt(d30)} − остатки {fmt(left3012)}; 30м №3 оценка {fmt(e303)} (экономия {fmt(N30-e303)})")
for c in ("сопоставимо", "единица/состав", "несопоставимо", "не использовано", "поза кошторисом", "—"):
    print(f"   30 категория {c:<16} Δ {fmt(T['d30|'+c])}")
print(f"ВЕРОЯТНО (8 домов): {fmt(probable)} = {fmt(probable/N9*100,1)} % нормы 8 домов ({fmt(N9)})")
print(f"Дефицит закупок по остаткам: 30м №3 {fmt(sum(r['e303_def'] for r in rows1))}, 30м №4–7 {fmt(sum(r['e47_def'] for r in rows1))}, итого {fmt(deficit)}")
print(f"EAC13 по фактическим остаткам = закупки {fmt(PUR)} + дефицит {fmt(deficit)} = {fmt(eac_fact13)}; экономия к СС13 {fmt(N13-eac_fact13)} ({fmt((N13-eac_fact13)/N13*100,1)} %)")
print(f"ETC отчёта − дефицит = {fmt(ETC-deficit)}")
for h, d in house.items():
    print(f"   {h}: отчёт норма {fmt(d['ПланГрн'])} факт {fmt(d['ФактВНормеГрн']+d['ФактПонадГрн'])} экономия {fmt(d['ЭкономияГрн'])} остаток {fmt(d['ОстатокОбъектаГрн'])} | реальн.брутто {fmt(cons_gt[h])} остаток {fmt(left_gt[h])}")

# ТОП-80 % по экономии 15 м (по модулю Δ среди экономий) и по отклонению отчёта директору
def top80(key, sign=-1):
    items = sorted([r for r in rows1 if sign * r[key] > 0], key=lambda r: -sign * r[key])
    tot = sum(sign * r[key] for r in items); acc = Z; out = []
    for r in items:
        acc += sign * r[key]; out.append((r, acc / tot * 100 if tot else Z))
        if acc >= tot * D("0.8"): break
    return out, tot
t15, tot15 = top80("d15_delta", -1)
tdir, totdir = top80("dev", -1)
tetc, totetc = top80("etc_minus_def", 1)
print("ТОП-80% экономии 15м:", len(t15), "назв из", sum(1 for r in rows1 if r['d15_delta'] < 0), "Σ экономий", fmt(tot15))
print("ТОП-80% экономии в отчёте директору:", len(tdir), "Σ", fmt(totdir))
print("ТОП-80% «ETC − дефицит»:", len(tetc), "Σ", fmt(totetc))

json.dump({"T": {k: str(v) for k, v in T.items()}, "LB": str(LB), "LB15": str(LB15), "LB30": str(LB30), "N8": str(N9), "left15": str(left15), "d15": str(d15), "d30": str(d30),
           "left3012": str(left3012), "e303": str(e303), "probable": str(probable), "deficit": str(deficit),
           "deficit303": str(sum(r['e303_def'] for r in rows1)), "deficit47": str(sum(r['e47_def'] for r in rows1)),
           "eac13": str(eac_fact13), "house": {h: {c: str(v) for c, v in d.items()} for h, d in house.items()},
           "cons_gt": {h: str(v) for h, v in cons_gt.items()}, "left_gt": {h: str(v) for h, v in left_gt.items()},
           "top15": [(r["name"], str(r["d15_delta"]), str(p)) for r, p in t15],
           "topdir": [(r["name"], str(r["dev"]), str(p)) for r, p in tdir],
           "topetc": [(r["name"], str(r["etc_minus_def"]), str(p)) for r, p in tetc]},
          open(os.path.join(OUT, "final.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=0)

# ---------- XLSX
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
wb = openpyxl.Workbook()
B = Font(bold=True); HDR = PatternFill("solid", fgColor="DDEBF7")

def sheet(title, header, data, widths=None, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True); r0 = 3
    for j, h in enumerate(header, 1):
        c = ws.cell(r0, j, h); c.font = B; c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(data, r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, Decimal):
                v = float(r2(v)) if abs(v) >= D("0.005") or v == 0 else float(v)
            c = ws.cell(i, j, v)
            if isinstance(v, float):
                c.number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    for j in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths[j - 1] if widths and j <= len(widths) else 14)
    return ws

wb.remove(wb.active)
u = lambda r, p: ", ".join(x.split(":", 1)[1].strip() for x in r["units"] if x.startswith(p)) or ""
P1H = ["Загальна назва", "Од. СС", "Од. закупки ERP", "Од. BuhBud",
       "Норма/дім 15 м, к-сть", "Ціна СС 15 м", "Норма/дім 15 м, грн", "Норма/дім 30 м, к-сть", "Ціна СС 30 м", "Норма/дім 30 м, грн",
       "Норма 8 будинків, к-сть", "Норма 8 будинків, грн", "Норма 13 будинків, к-сть", "Норма 13 будинків, грн",
       "Закуплено (ERP), к-сть", "Закуплено, грн з ПДВ", "Ціна закупки сер., з ПДВ", "Ціна закупки мін", "Ціна закупки макс",
       "Ціна економіста 15 м", "Обсяг економіста/дім 15 м", "Ціна економіста 30 м", "Обсяг економіста/дім 30 м",
       "Списано 15 м×6, к-сть", "Списано на дім: сер", "мін", "макс", "Списано 15 м×6, грн (реальн. брутто)",
       "Δ к-сть 15 м, %", "Δ ціна 15 м, %", "Δ 15 м, грн", "Ціновий ефект 15 м", "Кількісний ефект 15 м", "Категорія 15 м",
       "Будинків 15 м без списання", "Списано 30 м №1–2, к-сть", "Списано 30 м №1–2, грн", "Δ 30 м №1–2, грн", "Категорія 30 м",
       "Звіт директору: відхилення (розрах.)", "ціновий чинник", "кількісний чинник", "ETC (розрах.)",
       "Склад проекту, к-сть", "Склад проекту, грн брутто", "Залишок складів 8 домиків, к-сть", "Дефіцит до норми 30 м №3, грн", "Дефіцит до норми 30 м №4–7, грн",
       "ETC − дефіцит, грн", "Порівнюваність ціни (≥10×/≤0,1×)"]
data = []
for r in sorted(rows1, key=lambda r: r["d15_delta"]):
    o = ON[r["k"]]
    data.append([r["name"], u(r, "СС"), u(r, "ERP"), u(r, "BUH"),
                 r["n15_q"], r["p15"], r["n15_s"], r["n30_q"], r["p30"], r["n30_s"], r["n9_q"], r["n9_s"], r["n13_q"], r["n13_s"],
                 r["pu_q"], r["pu_g"], r["p_pu"], r["pu_pmin"], r["pu_pmax"],
                 r["econ_p15"], r["econ_q15"], r["econ_p30"], r["econ_q30"],
                 r["d15_cq"], r["d15_avg"], r["d15_min"], r["d15_max"], r["d15_cg"],
                 r["d15_dq_pct"], r["d15_dp_pct"], r["d15_delta"], r["d15_pe"], r["d15_qe"], r["c15"], r["d15_zero_houses"],
                 r["d30_cq"], r["d30_cg"], r["d30_delta"], r["c30"],
                 r["dev"], r["dir_pf"], r["dir_qf"], r["etc"],
                 o["stk|P|q"], o["stk|P|gt"], o["stk|H_scope|q"], r["e303_def"], r["e47_def"], r["etc_minus_def"],
                 "несопоставимо" if r["incomparable"] else ""])
tot_row = ["ИТОГО"] + [""] * 3
for j in range(4, len(P1H)):
    col = [x[j] for x in data]
    tot_row.append(sum((v for v in col if isinstance(v, Decimal)), Z) if j in (6, 9, 11, 13, 15, 27, 30, 31, 32, 36, 37, 39, 40, 41, 42, 44, 46, 47, 48) else "")
data.append(tot_row)
sheet("Ф1 по назвах", P1H, data, [34] + [9] * 3 + [12] * 46,
      "Расчёт по выгрузкам ERP/BuhBud на 02.10.2026. Норма — СС (ТЧ Комплектующие, на дом); закупки — ERP СебестоимостьТоваров (ветка 5 отчёта директору); "
      "списание — BuhBud Хозрасчетный Кт 20/22 (Комплектація, ПМА) по складам домиков, брутто по реальному НДС партии; эффекты: ціновий = к-сть × (ціна факт − ціна СС), кількісний = (к-сть факт − норма) × ціна СС.")
sheet("Ф1 ТОП-80%", ["Загальна назва", "Δ 15 м×6, грн", "Накопичено, %", "Категорія", "Ціновий", "Кількісний"],
      [[r["name"], r["d15_delta"], p, r["c15"], r["d15_pe"], r["d15_qe"]] for r, p in t15], [40, 16, 12, 18, 16, 16],
      f"Названия, дающие 80 % суммы экономий по расходу на 15 м №1–№6 (Σ экономий {fmt(tot15)} грн). Расчёт.")
sheet("Ф3 ETC−дефіцит ТОП", ["Загальна назва", "ETC − дефіцит, грн", "Накопичено, %", "ETC (розрах.)", "Дефіцит 30 м №3", "Дефіцит 30 м №4–7"],
      [[r["name"], r["etc_minus_def"], p, r["etc"], r["e303_def"], r["e47_def"]] for r, p in tetc], [40, 16, 12, 16, 16, 16],
      "Сколько «Плановий розрахунок» отчёта докупает по цене СС сверх дефицита, рассчитанного по фактическим остаткам (склад проекта + склады домиков). Расчёт.")

# мост (ф.2) — итоги в нетто BuhBud и реальном брутто
def S(fn):
    return sum((fn(ON[k]) for k in ON), Z)
mvh = lambda a, hs: S(lambda o: sum(o[f"mv|{a}>H:{h}|n"] for h in hs))
c_n = lambda hs: S(lambda o: sum(o[f"c|{h}|n"] for h in hs)); c_g = lambda hs: S(lambda o: sum(o[f"c|{h}|gt"] for h in hs))
s_n = lambda hs: S(lambda o: sum(o[f"stkh|{h}|n"] for h in hs)); s_g = lambda hs: S(lambda o: sum(o[f"stkh|{h}|gt"] for h in hs))
H3 = ["30-3"]
hh_back = S(lambda o: sum(o[m] for m in o if m.startswith("mv|H:") and ">P|n" in m))
inP = S(lambda o: sum(o[m] for m in o if m.startswith("in|") and m.endswith(">P|n")))
br = [
    ["1. Норма СС, 8 будинків у межах (з ПДВ)", "", N9, "СС: 704 612,01×6 + 1 009 864,74×2"],
    ["2. Закуплено (ERP), нетто / брутто", D("7634660.50"), PUR, "ERP СебестоимостьТоваров, Приход ПриобретениеТоваровУслуг ≤30.09.2026"],
    ["2а. Дод. витрати (лише BuhBud)", D("20169.17"), "", "BuhBud «Надходження дод. витрат» на склад проекту"],
    ["3. Надійшло на склад проекту (BuhBud)", inP, "", "ПТУ + авансові звіти + дод. витрати"],
    ["3а. Надійшло на склад МШП / чужий склад", S(lambda o: o["in|631>MSP|n"] + o["in|631>X|n"]), "", "ІБ_МШП 72 042,90; ІС Основной склад 8 680,00"],
    ["4. Склад проекту → склади 8 домиків", mvh("P", SC8), "", "ПеремещениеТоваров"],
    ["4а. Склад проекту → склад 30 м №3 (поза межами)", mvh("P", H3), "", "ПеремещениеТоваров"],
    ["4б. Отримано домиками (8) з чужих складів", mvh("X", SC8), "", "ІБ Загальний склад Виробництво, ІС Основной склад, ВООЗ, ООН"],
    ["4в. Отримано 30 м №3 з чужих складів", mvh("X", H3), "", ""],
    ["4г. Отримано домиками зі складу МШП (8 / №3)", mvh("MSP", SC8) + mvh("MSP", H3), "", f"з них 30 м №3: {fmt(mvh('MSP', H3))}"],
    ["4д. Повернуто з домиків на склад проекту", -hh_back, "", ""],
    ["5. Списано 8 будинків (Комплектація + ПМА)", -c_n(SC8), -c_g(SC8), "брутто — за реальним ПДВ партії"],
    ["5а. Списано 30 м №3 (не завершено, поза межами)", -c_n(H3), -c_g(H3), ""],
    ["6. Залишок на складах 8 домиків", s_n(SC8), s_g(SC8), "з них 15 м №4: 36 921,96; 30 м №2: ≈99,7 тис."],
    ["6а. Залишок на складі 30 м №3", s_n(H3), s_g(H3), "поза межами"],
    ["7. Залишок на складі проекту", S(lambda o: o["stk|P|n"]), S(lambda o: o["stk|P|gt"]), "ІБ _МД IRC 2026 та ін. склади проекту"],
    ["8. Передано зі складу проекту на інші проекти / МШП", -S(lambda o: o["mv|P>X|n"] + o["mv|P>MSP|n"]), "", ""],
    ["9. Списано зі складу проекту (паливо, Списання товарів)", -S(lambda o: sum(o[m] for m in o if m.startswith("out|P|") and m.endswith("|n"))), "", ""],
]
sheet("Ф2 міст", ["Крок", "Нетто (BuhBud), грн", "Брутто, грн", "Джерело"], br, [52, 18, 18, 70],
      "Мост «СС → закупка → склад домика → списание → остаток», IRC, все документы в базах на 02.10.2026. Границы: 15 м №1–№6, 30 м №1–№2. Невязка по нетто = 0 (проверено f32).")
hh = []
for h, d in house.items():
    rk = next((x for x in RK if (("15 м" in x["Спец"]) == h.startswith("15")) and x["Подр"].strip().endswith("№" + h.split("-")[1])), None)
    hh.append([h, (rk or {}).get("Номер", ""), (rk or {}).get("Статус", ""), (rk or {}).get("D", ""), (rk or {}).get("M", ""),
               "так" if (rk or {}).get("Подтв") else "ні", d["ПланГрн"], d["ФактВНормеГрн"], d["ФактПонадГрн"],
               d["ФактВНормеГрн"] + d["ФактПонадГрн"], d["ЭкономияГрн"], d["ОстатокОбъектаГрн"], cons_gt.get(h, Z), left_gt.get(h, Z)])
sheet("Ф2 будинки", ["Будинок", "РК №", "Статус РК", "D (відвант.)", "M (монтаж)", "Списання підтв.", "Норма, грн (звіт)",
                     "В нормі, грн (звіт)", "Понад, грн (звіт)", "Всього списано, грн (звіт)", "Економія, грн (звіт)", "Залишок домика, грн (звіт)",
                     "Списано, реальн. брутто (розрах.)", "Залишок домика, реальн. брутто (розрах.)"], hh, [9, 13, 20, 16, 16, 10] + [16] * 8,
      "Отчёт «Списання за нормами і понад норму (бухоблік)», ПолучитьДанные за 01.08–31.10.2026 (= «Списання IRC.xlsx»), + РасчетКомплектаций.")
mv = [[x["Контрагент"], D(x["Закуплено"]), D(x["Оплачено≤30.09"]), D(x["Сальдо(+долг/−аванс)"])] for x in MVP["lines"]]
sheet("Ф4 гроші", ["Контрагент", "Закуплено IRC (ERP), з ПДВ", "Оплачено з казни (матеріали+МШП) ≤30.09", "Сальдо: + борг / − аванс"], mv, [44, 18, 18, 18],
      "ERP: СебестоимостьТоваров (закупки IRC) vs РН.А_ДвиженияДенегИзКазны (статьи УТ-001067, УТ-001065; подразделение в иерархии МД IRC 2026). Расчёт.")
ev = []
for sh in ("IRS 15", "IRS 30"):
    for x in EV[sh]["rows"]:
        ev.append([sh, x["row"], x["name"], x["unit"], D(x["q"]), D(x["p"]), D(x["s"]), x["match"], x["ОН"], D(x["orig_q"]), x["orig_unit"], D(x["orig_s"]), D(x["ds"])])
sheet("Ф4 кошторис економіста", ["Лист", "Рядок", "Найменування (економіст)", "Од.", "Обсяг/дім", "Ціна з ПДВ", "Вартість/дім",
                                 "Збіг", "Загальна назва СС", "Обсяг оригіналу СС", "Од. оригіналу", "Сума оригіналу СС", "Δ економіст − СС"], ev,
      [8, 7, 44, 8, 10, 12, 14, 14, 26, 12, 10, 14, 14], "МАТЕРИАЛИ.xlsx против оригинала СС (BuhBud КомплектующиеПоМодульно). Расчёт.")
p = os.path.join(ROOT, r"_Rarzrabotki\notebook\specs")
os.makedirs(p, exist_ok=True)
xl = os.path.join(os.path.dirname(os.path.abspath(__file__)), r"..\..\..\notebook\specs\2026-10-02-ekonomiya-materialiv-ss-irc-ANALIZ.xlsx")
xl = os.path.abspath(xl)
wb.save(xl)
print("XLSX:", xl)
