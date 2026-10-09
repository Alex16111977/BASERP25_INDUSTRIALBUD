# -*- coding: utf-8 -*-
"""Дамп ячеек xlsx/xls (значения и формулы) в текст. Только чтение исходников."""
import sys, os, zipfile, io
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl, xlrd

HERE = os.path.dirname(os.path.abspath(__file__))


def fixed_copy(path):
    """1С кладёт xl/SharedStrings.xml с заглавной S — переписать zip в scratch."""
    out = os.path.join(HERE, "_fix_" + os.path.basename(path))
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            name = it.filename
            if name == "xl/SharedStrings.xml":
                name = "xl/sharedStrings.xml"
            if name == "[Content_Types].xml" or name.endswith(".rels"):
                data = data.replace(b"SharedStrings.xml", b"sharedStrings.xml")
            zo.writestr(name, data)
    return out


def load(path, data_only):
    try:
        return openpyxl.load_workbook(path, data_only=data_only)
    except KeyError:
        return openpyxl.load_workbook(fixed_copy(path), data_only=data_only)


def dump_xlsx(path, sheets=None):
    wf, wv = load(path, False), load(path, True)
    print("=====", path, wf.sheetnames)
    for ws in wf.worksheets:
        if sheets and ws.title not in sheets:
            continue
        wsv = wv[ws.title]
        print("--", ws.title, ws.dimensions)
        for row in ws.iter_rows():
            cells = []
            for c in row:
                if c.value is None:
                    continue
                v = wsv[c.coordinate].value
                if isinstance(c.value, str) and c.value.startswith("="):
                    cells.append("%s=%r {%s}" % (c.coordinate, v, c.value))
                else:
                    cells.append("%s=%r" % (c.coordinate, c.value))
            if cells:
                print(" | ".join(cells))


def dump_xls(path, sheets=None):
    wb = xlrd.open_workbook(path, formatting_info=False)
    print("=====", path, wb.sheet_names())
    for sh in wb.sheets():
        if sheets and sh.name not in sheets:
            continue
        print("--", sh.name, sh.nrows, sh.ncols)
        for r in range(sh.nrows):
            cells = []
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                if v not in ("", None):
                    cells.append("%s%d=%r" % (xlrd.formula.colname(c), r + 1, v))
            if cells:
                print(" | ".join(cells))


if __name__ == "__main__":
    p = sys.argv[1]
    sheets = sys.argv[2:] or None
    (dump_xls if p.lower().endswith(".xls") else dump_xlsx)(p, sheets)
