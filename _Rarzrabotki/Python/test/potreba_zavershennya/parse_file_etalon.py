# -*- coding: utf-8 -*-
"""Python-эталон разбора файла потребы директора (правила промта 2026-10-08, раздел 3.2).

Независимая реализация на openpyxl: результат загрузчика 1С (ПрочитатьФайлПотребы) обязан совпасть.
    python parse_file_etalon.py <файл.xlsx> [--json out.json]
"""
import json
import re
import shutil
import sys
import tempfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

APOS = "'’ʼ‘´`"


def norm(s):
    s = "" if s is None else str(s)
    s = s.replace(" ", " ")
    for a in APOS:
        s = s.replace(a, "'")
    s = re.sub(r"\s+", " ", s).strip().lower()
    return s.replace("i", "і")


def norm_head(s):
    return norm(s).replace("'", "").replace("-", "")


def r2(x):
    return float(Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def field_of(h):
    """Поле по нормализованному заголовку (None — не распознано)."""
    ss = "сс" in h.split() or h.endswith(" сс") or " сс " in h
    if not h:
        return None
    if "вірна назва" in h or "загальна назва" in h or "назва згідно сс" in h or "назва сс" in h:
        return "НазваСС"
    if "назва товару" in h or "найменування" in h or h == "товар":
        return "НазваТовару"
    if ss and ("обєм" in h or "обсяг" in h or "кількість" in h or "ксть" in h):
        return "ОбсягСС"
    if ss and ("од" in h.split() or h.startswith("од ") or "одиниц" in h):
        return "ОдСС"
    if not ss and ("одиниц" in h or "од вим" in h):
        return "Од"
    if not ss and ("ксть" in h or "кількість" in h or "кть" in h or "колво" in h):
        return "КСть"
    if "будин" in h or "дім" in h:
        return "Будинок"
    if "ціна" in h:
        return "Ціна"
    if "вартість" in h or "сума" in h:
        return "Вартість"
    return None


def detect_header(row_vals):
    """{поле: колонка} или None."""
    found = {}
    od_cols = []
    for c, v in enumerate(row_vals, start=1):
        f = field_of(norm_head(v))
        if f is None:
            continue
        if f == "Од":
            od_cols.append(c)
            continue
        found.setdefault(f, c)
    if od_cols:
        kst = found.get("КСть")
        left = [c for c in od_cols if kst and c < kst]
        found["Од"] = max(left) if left else od_cols[0]
    ok = (("НазваСС" in found or "НазваТовару" in found)
          and ("ОбсягСС" in found or "КСть" in found)
          and ("Вартість" in found or "Ціна" in found))
    return found if ok and len(found) >= 4 else None


def houses_text(cell):
    v = cell.value
    if v is None:
        return ""
    if is_num(v):
        # как 1С .Текст: формат ячейки; дробь 6.7 Excel показывает «6,7»
        t = (str(int(v)) if float(v).is_integer() else str(v))
    else:
        t = str(v)
    t = re.sub(r"[.;\s]+", ",", t.replace(",", ","))
    t = re.sub(r",+", ",", t).strip(",")
    return t


def parse(path):
    tmp = Path(tempfile.mkdtemp()) / "potreba.xlsx"
    shutil.copy(path, tmp)
    wb = openpyxl.load_workbook(tmp, data_only=True)
    sections = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=30))
        vals = [[c.value for c in r] for r in rows]

        def is_marker(i):
            nz = [v for v in vals[i] if v not in (None, "")]
            return len(nz) == 1 and norm_head(nz[0]).startswith("проект")

        markers = [i for i in range(len(vals)) if is_marker(i)]
        starts = markers or [None]
        for k, m in enumerate(starts):
            if m is None:
                name, lo, hi = ws.title, 0, len(vals)
                hdr_range = range(0, min(20, len(vals)))
            else:
                name = str([v for v in vals[m] if v not in (None, "")][0]).strip()
                lo = m + 1
                hi = markers[k + 1] if k + 1 < len(markers) else len(vals)
                hdr_range = range(lo, min(lo + 3, hi))
            hdr, hdr_i = None, None
            for i in hdr_range:
                hdr = detect_header(vals[i])
                if hdr:
                    hdr_i = i
                    break
            if not hdr:
                continue

            def g(i, f):
                c = hdr.get(f)
                return vals[i][c - 1] if c else None

            data, itogo, empty_run = [], None, 0
            for i in range(hdr_i + 1, hi):
                a, b = g(i, "НазваТовару"), g(i, "НазваСС")
                nums = [g(i, f) for f in ("ОбсягСС", "КСть", "Ціна", "Вартість")]
                if all(v in (None, "") for v in vals[i]):
                    empty_run += 1
                    if empty_run >= 3:
                        break
                    continue
                empty_run = 0
                rowtext = " ".join(norm(v) for v in vals[i] if isinstance(v, str))
                if (norm(a) == "" and norm(b) == "" and is_num(g(i, "Вартість"))) or \
                        any(w in rowtext for w in ("разом", "итого", "всього", "усього")):
                    itogo = g(i, "Вартість")
                    break
                if (norm(a) or norm(b)) and any(is_num(v) for v in nums):
                    J, G, I, E = g(i, "Вартість"), g(i, "КСть"), g(i, "Ціна"), g(i, "ОбсягСС")
                    warn = []
                    suma = r2(J) if is_num(J) else (r2(G * I) if is_num(G) and is_num(I) else 0.0)
                    if is_num(J) and is_num(G) and is_num(I) and abs(J - G * I) > 0.01:
                        warn.append("Вартість ≠ К-сть × Ціна")
                    cena = r2(I) if is_num(I) else (r2(suma / G) if is_num(G) and G else 0.0)
                    if is_num(E) and E > 0:
                        kss = round(float(E), 3)
                    elif is_num(G) and norm(g(i, "ОдСС")) == norm(g(i, "Од")):
                        kss = round(float(G), 3)
                    else:
                        kss = 0.0
                        warn.append("Не визначено кількість в од. СС")
                    nss = str(b).strip() if b is not None else ""
                    if "/" in nss:
                        warn.append("Складена назва")
                    hcell = ws.cell(row=i + 1, column=hdr["Будинок"]) if "Будинок" in hdr else None
                    data.append({
                        "row": i + 1, "НазваниеТовара": str(a).strip() if a else "", "НазваниеСС": nss,
                        "Количество": round(float(G), 3) if is_num(G) else 0.0,
                        "Единица": str(g(i, "Од") or "").strip(), "КоличествоСС": kss,
                        "Цена": cena, "Сумма": suma,
                        "Дома": houses_text(hcell) if hcell else "", "warn": warn,
                    })
            if data:
                sections.append({"name": name, "sheet": ws.title, "rows": data,
                                 "sum": r2(sum(d["Сумма"] for d in data)),
                                 "itogo": r2(itogo) if is_num(itogo) else None})
    return sections


if __name__ == "__main__":
    src = sys.argv[1]
    secs = parse(src)
    for s in secs:
        warn = sum(1 for d in s["rows"] if d["warn"])
        print(f"{s['name']!r} [{s['sheet']}]: строк {len(s['rows'])}, Σ {s['sum']:.2f}, ИТОГО {s['itogo']}, с предупреждениями {warn}")
        for d in s["rows"]:
            if d["warn"]:
                print(f"    r{d['row']}: {d['НазваниеСС']} — {'; '.join(d['warn'])}")
    if "--json" in sys.argv:
        Path(sys.argv[sys.argv.index("--json") + 1]).write_text(json.dumps(secs, ensure_ascii=False, indent=1), encoding="utf-8")
